"""Work/funding cache management views and exports."""

import csv
import logging
from io import BytesIO, StringIO

from babel import Locale, UnknownLocaleError
from flask import Response, request, send_file, stream_with_context
from flask_babel import _, get_locale
from sqlalchemy import and_, case, func, or_
from sqlalchemy.orm import load_only

from .. import db, plain_text
from .works_shared import (
    OPENALEX_EXPORT_CSV_COLUMNS,
    OPENALEX_EXPORT_XLSX_COLUMNS,
    _attach_download_token,
    _excel_cell,
    _openalex_cache_key_expr,
    _openalex_export_load_options,
    _openalex_export_row,
    _openalex_normalized_doi_expr,
    _pagination_dict,
    _table_page_params,
    _table_sort_params,
)
logger = logging.getLogger(__name__)
from .works_state import _OPENALEX_OA_COLORS, _OPENALEX_PRIORITY_OA_COLORS
def _send_openalex_export(records_query, base_name: str = 'openalex_articles_all_institutions'):
    """Send OpenAlex exports without materializing the full dataset as a DataFrame."""
    export_format = (request.args.get('format') or '').lower()

    if export_format == 'excel':
        from openpyxl import Workbook

        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet('OpenAlex')
        worksheet.append(OPENALEX_EXPORT_XLSX_COLUMNS)
        for raw, metadata in records_query.yield_per(500):
            row = _openalex_export_row(raw, metadata, include_raw_json=False)
            worksheet.append([_excel_cell(row.get(column)) for column in OPENALEX_EXPORT_XLSX_COLUMNS])

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return _attach_download_token(send_file(
            output,
            as_attachment=True,
            download_name=f"{base_name}.xlsx",
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        ))

    def generate_csv():
        buffer = StringIO()
        writer = csv.DictWriter(buffer, fieldnames=OPENALEX_EXPORT_CSV_COLUMNS, extrasaction='ignore')
        buffer.write('\ufeff')
        writer.writeheader()
        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)

        for raw, metadata in records_query.yield_per(500):
            writer.writerow(_openalex_export_row(raw, metadata))
            yield buffer.getvalue()
            buffer.seek(0)
            buffer.truncate(0)

    response = Response(
        stream_with_context(generate_csv()),
        mimetype='text/csv; charset=utf-8',
    )
    response.headers['Content-Disposition'] = f'attachment; filename="{base_name}.csv"'
    return _attach_download_token(response)


def _openalex_fact_cache_summary(ror_id: str) -> dict:
    from ..models import (
        OpenAlexInstitutionWorkFact,
        OpenAlexSyncRun,
    )
    from ..services.openalex_service import TITLE_MATCH_NOT_FOUND_ERROR

    fact = OpenAlexInstitutionWorkFact
    title_retryable = or_(
        fact.raw_error.is_(None),
        fact.raw_error != TITLE_MATCH_NOT_FOUND_ERROR,
    )
    title_candidate = and_(
        fact.has_local_title.is_(True),
        title_retryable,
        or_(
            and_(
                fact.has_valid_doi.is_(False),
                or_(
                    fact.raw_status.is_(None),
                    fact.raw_status != "found",
                ),
            ),
            and_(
                fact.has_valid_doi.is_(True),
                fact.raw_status == "not_found",
                fact.openalex_id.is_(None),
            ),
        ),
    )
    row = (
        db.session.query(
            func.coalesce(func.sum(fact.source_record_count), 0),
            func.coalesce(func.sum(case(
                (fact.has_valid_doi.is_(True), fact.source_record_count),
                else_=0,
            )), 0),
            func.count(case((fact.has_valid_doi.is_(True), 1))),
            func.count(case((
                and_(
                    fact.has_valid_doi.is_(True),
                    fact.raw_status.isnot(None),
                ),
                1,
            ))),
            func.count(case((
                and_(
                    fact.has_valid_doi.is_(True),
                    fact.openalex_id.isnot(None),
                ),
                1,
            ))),
            func.count(case((
                and_(
                    fact.has_valid_doi.is_(True),
                    fact.raw_status == "not_found",
                ),
                1,
            ))),
            func.count(case((
                and_(
                    fact.has_valid_doi.is_(True),
                    fact.raw_status == "error",
                ),
                1,
            ))),
            func.count(case((fact.openalex_id.isnot(None), 1))),
            func.coalesce(func.sum(case(
                (title_candidate, fact.source_record_count),
                else_=0,
            )), 0),
            func.coalesce(func.sum(case((
                and_(
                    title_candidate,
                    fact.has_valid_doi.is_(False),
                ),
                fact.source_record_count,
            ), else_=0)), 0),
            func.coalesce(func.sum(case((
                and_(
                    title_candidate,
                    fact.has_valid_doi.is_(True),
                ),
                fact.source_record_count,
            ), else_=0)), 0),
        )
        .filter(fact.ror_id == ror_id)
        .one()
    )
    article_works = int(row[0] or 0)
    article_doi_works = int(row[1] or 0)
    candidate_dois = int(row[2] or 0)
    processed = int(row[3] or 0)
    matched = int(row[4] or 0)
    not_found = int(row[5] or 0)
    errors = int(row[6] or 0)
    matched_cache_keys = int(row[7] or 0)
    no_doi_title_candidates = int(row[9] or 0)
    doi_not_found_title_candidates = int(row[10] or 0)
    last_run = (
        OpenAlexSyncRun.query
        .filter_by(ror_id=ror_id)
        .order_by(OpenAlexSyncRun.finished_at.desc())
        .first()
    )
    return {
        "article_works": article_works,
        "article_doi_works": article_doi_works,
        "candidate_dois": candidate_dois,
        "processed_dois": processed,
        "matched_dois": matched,
        "matched_openalex_keys": matched_cache_keys,
        "not_found_dois": not_found,
        "error_dois": errors,
        "pending_dois": max(candidate_dois - processed, 0),
        "unmatched_dois": max(
            candidate_dois - matched - max(candidate_dois - processed, 0) - errors,
            0,
        ),
        "title_candidate_works": int(row[8] or 0),
        "no_doi_title_candidates": no_doi_title_candidates,
        "doi_not_found_title_candidates": doi_not_found_title_candidates,
        "processed_percent": round((processed / candidate_dois * 100), 1) if candidate_dois else 0,
        "matched_percent": round((matched / candidate_dois * 100), 1) if candidate_dois else 0,
        "last_run": last_run,
    }


def _openalex_cache_summary(ror_id: str) -> dict:
    """Summarize OpenAlex sync status for DOI-backed journal articles."""
    from ..models import OpenAlexSyncRun, OpenAlexWorkMetadata, OpenAlexWorkRawCache, WorkCache
    from ..services.analytics_service import openalex_fact_available
    from ..services.openalex_service import TITLE_MATCH_NOT_FOUND_ERROR

    if openalex_fact_available(ror_id):
        return _openalex_fact_cache_summary(ror_id)

    normalized = _openalex_normalized_doi_expr(WorkCache).label("doi_normalized")
    cache_key = _openalex_cache_key_expr(WorkCache).label("openalex_cache_key")
    article_filters = (
        WorkCache.ror_id == ror_id,
        WorkCache.type == "journal-article",
    )
    doi_filters = (
        WorkCache.ror_id == ror_id,
        WorkCache.doi_normalized.isnot(None),
        WorkCache.type == "journal-article",
    )
    doi_subquery = (
        db.session.query(normalized)
        .filter(*doi_filters)
        .distinct()
        .subquery()
    )
    cache_key_subquery = (
        db.session.query(cache_key)
        .filter(*article_filters)
        .distinct()
        .subquery()
    )

    article_works = db.session.query(func.count(WorkCache.id)).filter(*article_filters).scalar() or 0
    article_doi_works = db.session.query(func.count(WorkCache.id)).filter(*doi_filters).scalar() or 0
    candidate_dois = db.session.query(func.count()).select_from(doi_subquery).scalar() or 0
    processed = (
        db.session.query(func.count())
        .select_from(doi_subquery)
        .join(OpenAlexWorkRawCache, OpenAlexWorkRawCache.doi_normalized == doi_subquery.c.doi_normalized)
        .scalar()
        or 0
    )
    matched = (
        db.session.query(func.count())
        .select_from(doi_subquery)
        .join(OpenAlexWorkMetadata, OpenAlexWorkMetadata.doi_normalized == doi_subquery.c.doi_normalized)
        .scalar()
        or 0
    )
    not_found = (
        db.session.query(func.count())
        .select_from(doi_subquery)
        .join(OpenAlexWorkRawCache, OpenAlexWorkRawCache.doi_normalized == doi_subquery.c.doi_normalized)
        .filter(OpenAlexWorkRawCache.status == "not_found")
        .scalar()
        or 0
    )
    errors = (
        db.session.query(func.count())
        .select_from(doi_subquery)
        .join(OpenAlexWorkRawCache, OpenAlexWorkRawCache.doi_normalized == doi_subquery.c.doi_normalized)
        .filter(OpenAlexWorkRawCache.status == "error")
        .scalar()
        or 0
    )
    matched_cache_keys = (
        db.session.query(func.count())
        .select_from(cache_key_subquery)
        .join(OpenAlexWorkMetadata, OpenAlexWorkMetadata.doi_normalized == cache_key_subquery.c.openalex_cache_key)
        .scalar()
        or 0
    )
    no_doi_title_candidates = (
        db.session.query(func.count())
        .select_from(WorkCache)
        .outerjoin(OpenAlexWorkRawCache, OpenAlexWorkRawCache.doi_normalized == cache_key)
        .filter(
            *article_filters,
            WorkCache.title.isnot(None),
            WorkCache.title != "",
            WorkCache.doi_normalized.is_(None),
            (OpenAlexWorkRawCache.id.is_(None)) |
            (
                (OpenAlexWorkRawCache.status != "found") &
                (
                    (OpenAlexWorkRawCache.error.is_(None)) |
                    (OpenAlexWorkRawCache.error != TITLE_MATCH_NOT_FOUND_ERROR)
                )
            ),
        )
        .scalar()
        or 0
    )
    doi_not_found_title_subquery = (
        db.session.query(normalized)
        .select_from(WorkCache)
        .join(OpenAlexWorkRawCache, OpenAlexWorkRawCache.doi_normalized == normalized)
        .outerjoin(OpenAlexWorkMetadata, OpenAlexWorkMetadata.doi_normalized == normalized)
        .filter(
            *doi_filters,
            WorkCache.title.isnot(None),
            WorkCache.title != "",
            OpenAlexWorkRawCache.status == "not_found",
            OpenAlexWorkMetadata.id.is_(None),
            (OpenAlexWorkRawCache.error.is_(None)) |
            (OpenAlexWorkRawCache.error != TITLE_MATCH_NOT_FOUND_ERROR),
        )
        .distinct()
        .subquery()
    )
    doi_not_found_title_candidates = (
        db.session.query(func.count())
        .select_from(doi_not_found_title_subquery)
        .scalar()
        or 0
    )
    last_run = (
        OpenAlexSyncRun.query
        .filter_by(ror_id=ror_id)
        .order_by(OpenAlexSyncRun.finished_at.desc())
        .first()
    )

    return {
        "article_works": article_works,
        "article_doi_works": article_doi_works,
        "candidate_dois": candidate_dois,
        "processed_dois": processed,
        "matched_dois": matched,
        "matched_openalex_keys": matched_cache_keys,
        "not_found_dois": not_found,
        "error_dois": errors,
        "pending_dois": max(candidate_dois - processed, 0),
        "unmatched_dois": max(candidate_dois - matched - max(candidate_dois - processed, 0) - errors, 0),
        "title_candidate_works": no_doi_title_candidates + doi_not_found_title_candidates,
        "no_doi_title_candidates": no_doi_title_candidates,
        "doi_not_found_title_candidates": doi_not_found_title_candidates,
        "processed_percent": round((processed / candidate_dois * 100), 1) if candidate_dois else 0,
        "matched_percent": round((matched / candidate_dois * 100), 1) if candidate_dois else 0,
        "last_run": last_run,
    }


def _openalex_work_rows(
    ror_id: str,
    coverage: str = "all",
    page: int = 1,
    per_page: int = 25,
    search: str = "",
    sort: str = "citations",
    direction: str = "desc",
) -> tuple[list[dict], dict, dict]:
    """Join local ORCID articles with OpenAlex metadata by DOI or local work key."""
    from ..models import OpenAlexWorkMetadata, OpenAlexWorkRawCache, WorkCache

    summary = _openalex_cache_summary(ror_id)
    normalized = _openalex_normalized_doi_expr(WorkCache).label("doi_normalized")
    cache_key = _openalex_cache_key_expr(WorkCache).label("openalex_cache_key")
    query = (
        db.session.query(
            WorkCache.id.label("work_cache_id"),
            WorkCache.title,
            WorkCache.orcid,
            WorkCache.type,
            WorkCache.pub_year,
            WorkCache.journal_title,
            WorkCache.doi,
            normalized,
            cache_key,
            OpenAlexWorkRawCache.status.label("raw_status"),
            OpenAlexWorkRawCache.error.label("raw_error"),
            OpenAlexWorkMetadata.openalex_id,
            OpenAlexWorkMetadata.cited_by_count,
            OpenAlexWorkMetadata.is_oa,
            OpenAlexWorkMetadata.oa_status,
            OpenAlexWorkMetadata.source_name,
            OpenAlexWorkMetadata.source_issn_l,
            OpenAlexWorkMetadata.primary_topic_field,
            OpenAlexWorkMetadata.primary_topic_domain,
        )
        .outerjoin(OpenAlexWorkRawCache, OpenAlexWorkRawCache.doi_normalized == cache_key)
        .outerjoin(OpenAlexWorkMetadata, OpenAlexWorkMetadata.doi_normalized == cache_key)
        .filter(
            WorkCache.ror_id == ror_id,
            WorkCache.type == "journal-article",
        )
    )

    if coverage == "enriched":
        query = query.filter(OpenAlexWorkMetadata.id.isnot(None))
    elif coverage == "missing":
        query = query.filter(
            WorkCache.doi_normalized.isnot(None),
            OpenAlexWorkRawCache.id.is_(None),
        )
    elif coverage == "not_found":
        query = query.filter(OpenAlexWorkRawCache.status.in_(("not_found", "error")))
    elif coverage == "no_doi":
        query = query.filter(WorkCache.doi_normalized.is_(None))

    search = (search or "").strip()
    if search:
        pattern = f"%{search}%"
        query = query.filter(or_(
            WorkCache.title.ilike(pattern),
            WorkCache.doi.ilike(pattern),
            WorkCache.orcid.ilike(pattern),
            WorkCache.journal_title.ilike(pattern),
            OpenAlexWorkMetadata.source_name.ilike(pattern),
            OpenAlexWorkMetadata.primary_topic_field.ilike(pattern),
            OpenAlexWorkMetadata.primary_topic_domain.ilike(pattern),
            OpenAlexWorkMetadata.openalex_id.ilike(pattern),
        ))

    total_rows = query.count()
    pagination = _pagination_dict(page, per_page, total_rows)
    page = pagination["page"]

    status_sort = case(
        (OpenAlexWorkMetadata.id.isnot(None), 0),
        (OpenAlexWorkRawCache.status == "error", 3),
        (OpenAlexWorkRawCache.status == "not_found", 2),
        else_=1,
    )
    sort_columns = {
        "title": WorkCache.title,
        "year": WorkCache.pub_year,
        "citations": OpenAlexWorkMetadata.cited_by_count,
        "open_access": OpenAlexWorkMetadata.is_oa,
        "source": OpenAlexWorkMetadata.source_name,
        "status": status_sort,
    }
    sort = sort if sort in sort_columns else "citations"
    direction = direction if direction in {"asc", "desc"} else "desc"
    sort_column = sort_columns[sort]
    null_rank = case((sort_column.is_(None), 1), else_=0)
    primary_order = sort_column.asc() if direction == "asc" else sort_column.desc()
    query = query.order_by(null_rank.asc(), primary_order, WorkCache.title.asc(), WorkCache.id.asc())
    result_rows = query.offset((page - 1) * per_page).limit(per_page).all()

    rows = []
    for row in result_rows:
        raw_status = row.raw_status or "pending"
        openalex_id = row.openalex_id
        has_doi = bool(row.doi_normalized)
        if openalex_id:
            status_key = "matched"
        elif raw_status == "not_found":
            status_key = "not_found"
        elif raw_status == "error":
            status_key = "error"
        elif not has_doi:
            status_key = "no_doi"
        else:
            status_key = "pending"
        rows.append({
            "work_cache_id": row.work_cache_id,
            "title": plain_text(row.title),
            "orcid": row.orcid,
            "type": row.type or "",
            "pub_year": row.pub_year or "",
            "journal_title": row.journal_title or "",
            "doi": row.doi or "",
            "doi_normalized": row.doi_normalized,
            "openalex_cache_key": row.openalex_cache_key,
            "has_doi": has_doi,
            "matched": bool(openalex_id),
            "status_key": status_key,
            "raw_status": raw_status,
            "raw_error": row.raw_error or "",
            "openalex_id": openalex_id or "",
            "openalex_url": f"https://openalex.org/{openalex_id}" if openalex_id else "",
            "cited_by_count": row.cited_by_count,
            "is_oa": bool(row.is_oa) if openalex_id else None,
            "oa_status": row.oa_status or "",
            "source_name": row.source_name or "",
            "source_issn_l": row.source_issn_l or "",
            "primary_topic_field": row.primary_topic_field or "",
            "primary_topic_domain": row.primary_topic_domain or "",
        })

    has_doi_condition = WorkCache.doi_normalized.isnot(None)
    coverage_row = (
        db.session.query(
            func.count(WorkCache.id),
            func.coalesce(func.sum(case((OpenAlexWorkMetadata.id.isnot(None), 1), else_=0)), 0),
            func.coalesce(func.sum(case((and_(has_doi_condition, OpenAlexWorkRawCache.id.is_(None)), 1), else_=0)), 0),
            func.coalesce(func.sum(case((OpenAlexWorkRawCache.status.in_(("not_found", "error")), 1), else_=0)), 0),
            func.coalesce(func.sum(case((~has_doi_condition, 1), else_=0)), 0),
        )
        .select_from(WorkCache)
        .outerjoin(OpenAlexWorkRawCache, OpenAlexWorkRawCache.doi_normalized == _openalex_cache_key_expr(WorkCache))
        .outerjoin(OpenAlexWorkMetadata, OpenAlexWorkMetadata.doi_normalized == _openalex_cache_key_expr(WorkCache))
        .filter(WorkCache.ror_id == ror_id, WorkCache.type == "journal-article")
        .one()
    )

    summary.update({
        "total_article_works": summary["article_works"],
        "total_doi_works": summary["article_doi_works"],
        "unique_dois": summary["candidate_dois"],
        "enriched_unique_dois": summary["matched_openalex_keys"],
        "processed_unique_dois": summary["processed_dois"],
        "coverage_percent": round((coverage_row[1] / coverage_row[0] * 100), 1) if coverage_row[0] else 0,
        "total_rows": total_rows,
        "total_citations": (
            db.session.query(func.coalesce(func.sum(OpenAlexWorkMetadata.cited_by_count), 0))
            .filter(OpenAlexWorkMetadata.doi_normalized.in_(
                db.session.query(_openalex_cache_key_expr(WorkCache))
                .filter(
                    WorkCache.ror_id == ror_id,
                    WorkCache.type == "journal-article",
                )
                .distinct()
            ))
            .scalar()
            or 0
        ),
        "open_access_works": (
            db.session.query(func.count())
            .filter(OpenAlexWorkMetadata.doi_normalized.in_(
                db.session.query(_openalex_cache_key_expr(WorkCache))
                .filter(
                    WorkCache.ror_id == ror_id,
                    WorkCache.type == "journal-article",
                )
                .distinct()
            ))
            .filter(OpenAlexWorkMetadata.is_oa.is_(True))
            .scalar()
            or 0
        ),
        "coverage_counts": {
            "all": int(coverage_row[0] or 0),
            "enriched": int(coverage_row[1] or 0),
            "missing": int(coverage_row[2] or 0),
            "not_found": int(coverage_row[3] or 0),
            "no_doi": int(coverage_row[4] or 0),
        },
    })

    return rows, summary, pagination


OPENALEX_INSTITUTION_EXPORT_COLUMNS = [
    'institution_ror_id',
    'local_work_id',
    'orcid',
    'local_title',
    'local_type',
    'local_publication_year',
    'local_journal',
    'local_doi',
    'openalex_cache_key',
    'matched',
] + OPENALEX_EXPORT_XLSX_COLUMNS


def _openalex_institution_export_query(
    ror_id: str,
    coverage: str = "all",
    search: str = "",
    sort: str = "citations",
    direction: str = "desc",
):
    """Build an unbounded, low-memory OpenAlex export query for one institution."""
    from ..models import OpenAlexWorkMetadata, OpenAlexWorkRawCache, WorkCache

    cache_key = _openalex_cache_key_expr(WorkCache)
    query = (
        db.session.query(WorkCache, OpenAlexWorkRawCache, OpenAlexWorkMetadata)
        .outerjoin(OpenAlexWorkRawCache, OpenAlexWorkRawCache.doi_normalized == cache_key)
        .outerjoin(OpenAlexWorkMetadata, OpenAlexWorkMetadata.doi_normalized == cache_key)
        .filter(
            WorkCache.ror_id == ror_id,
            WorkCache.type == "journal-article",
        )
    )

    if coverage == "enriched":
        query = query.filter(OpenAlexWorkMetadata.id.isnot(None))
    elif coverage == "missing":
        query = query.filter(
            WorkCache.doi_normalized.isnot(None),
            OpenAlexWorkRawCache.id.is_(None),
        )
    elif coverage == "not_found":
        query = query.filter(OpenAlexWorkRawCache.status.in_(("not_found", "error")))
    elif coverage == "no_doi":
        query = query.filter(WorkCache.doi_normalized.is_(None))

    search = (search or "").strip()
    if search:
        pattern = f"%{search}%"
        query = query.filter(or_(
            WorkCache.title.ilike(pattern),
            WorkCache.doi.ilike(pattern),
            WorkCache.orcid.ilike(pattern),
            WorkCache.journal_title.ilike(pattern),
            OpenAlexWorkMetadata.source_name.ilike(pattern),
            OpenAlexWorkMetadata.primary_topic_field.ilike(pattern),
            OpenAlexWorkMetadata.primary_topic_domain.ilike(pattern),
            OpenAlexWorkMetadata.openalex_id.ilike(pattern),
        ))

    status_sort = case(
        (OpenAlexWorkMetadata.id.isnot(None), 0),
        (OpenAlexWorkRawCache.status == "error", 3),
        (OpenAlexWorkRawCache.status == "not_found", 2),
        else_=1,
    )
    sort_columns = {
        "title": WorkCache.title,
        "year": WorkCache.pub_year,
        "citations": OpenAlexWorkMetadata.cited_by_count,
        "open_access": OpenAlexWorkMetadata.is_oa,
        "source": OpenAlexWorkMetadata.source_name,
        "status": status_sort,
    }
    sort_column = sort_columns.get(sort, OpenAlexWorkMetadata.cited_by_count)
    direction = direction if direction in {"asc", "desc"} else "desc"
    null_rank = case((sort_column.is_(None), 1), else_=0)
    primary_order = sort_column.asc() if direction == "asc" else sort_column.desc()

    return query.options(
        load_only(
            WorkCache.id,
            WorkCache.ror_id,
            WorkCache.orcid,
            WorkCache.title,
            WorkCache.type,
            WorkCache.pub_year,
            WorkCache.journal_title,
            WorkCache.doi,
            WorkCache.doi_normalized,
        ),
        *_openalex_export_load_options(OpenAlexWorkRawCache, OpenAlexWorkMetadata),
    ).order_by(
        null_rank.asc(),
        primary_order,
        WorkCache.title.asc(),
        WorkCache.id.asc(),
    )


def _openalex_institution_export_row(work, raw, metadata) -> dict:
    """Combine local institutional context with flattened OpenAlex metadata."""
    cache_key = work.doi_normalized or f"work:{work.id}"
    row = {
        'institution_ror_id': work.ror_id,
        'local_work_id': work.id,
        'orcid': work.orcid,
        'local_title': plain_text(work.title),
        'local_type': work.type,
        'local_publication_year': work.pub_year,
        'local_journal': plain_text(work.journal_title),
        'local_doi': work.doi,
        'openalex_cache_key': cache_key,
        'matched': bool(metadata and metadata.openalex_id),
    }
    row.update(
        _openalex_export_row(
            raw,
            metadata,
            include_raw_json=False,
            cache_key=cache_key,
        )
    )
    return row


def _send_openalex_institution_export(
    records_query,
    ror_id: str,
    coverage: str = "all",
):
    """Stream CSV or build write-only XLSX rows for an institutional export."""
    export_format = (request.args.get('format') or '').lower()
    base_name = f"openalex_works_{ror_id}_{coverage}"

    if export_format == 'excel':
        from openpyxl import Workbook

        workbook = Workbook(write_only=True)
        sheet_number = 1
        worksheet = workbook.create_sheet('OpenAlex works')
        worksheet.append(OPENALEX_INSTITUTION_EXPORT_COLUMNS)
        worksheet_rows = 1
        for work, raw, metadata in records_query.yield_per(1000):
            if worksheet_rows >= 1_048_576:
                sheet_number += 1
                worksheet = workbook.create_sheet(f'OpenAlex works {sheet_number}')
                worksheet.append(OPENALEX_INSTITUTION_EXPORT_COLUMNS)
                worksheet_rows = 1
            row = _openalex_institution_export_row(work, raw, metadata)
            worksheet.append([
                _excel_cell(row.get(column))
                for column in OPENALEX_INSTITUTION_EXPORT_COLUMNS
            ])
            worksheet_rows += 1

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return _attach_download_token(send_file(
            output,
            as_attachment=True,
            download_name=f"{base_name}.xlsx",
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        ))

    def generate_csv():
        buffer = StringIO()
        writer = csv.DictWriter(
            buffer,
            fieldnames=OPENALEX_INSTITUTION_EXPORT_COLUMNS,
            extrasaction='ignore',
        )
        buffer.write('\ufeff')
        writer.writeheader()
        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)

        for work, raw, metadata in records_query.yield_per(1000):
            writer.writerow(_openalex_institution_export_row(work, raw, metadata))
            yield buffer.getvalue()
            buffer.seek(0)
            buffer.truncate(0)

    response = Response(
        stream_with_context(generate_csv()),
        mimetype='text/csv; charset=utf-8',
    )
    response.headers['Content-Disposition'] = f'attachment; filename="{base_name}.csv"'
    return _attach_download_token(response)


def _int_filter(value) -> int | None:
    try:
        return int(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _list_filter(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        values = [value]
    else:
        values = list(value)
    return [str(item).strip() for item in values if str(item).strip()]


def _request_list_arg(name: str) -> list[str]:
    values = request.args.getlist(name)
    if values:
        return _list_filter(values)
    return _list_filter(request.args.get(name))


def _selected_values(selected: list[str], allowed: list[str], default: list[str]) -> list[str]:
    allowed_set = set(allowed)
    values = [value for value in selected if value in allowed_set]
    return values or default


def _is_export_request() -> bool:
    return request.endpoint in {"works.openalex_global_export", "works.openalex_analytics_export", "works.openalex_works_export"}


def _chart_color(index: int) -> str:
    colors = [
        "#2a69b8", "#28a745", "#f39c12", "#dc3545", "#17a2b8",
        "#6f42c1", "#20c997", "#6c757d", "#fd7e14", "#6610f2",
    ]
    return colors[index % len(colors)]


def _openalex_oa_status_label(value: str | None) -> str:
    """Return a localized, reader-facing label for an OpenAlex OA status."""
    status = str(value or "").strip().lower()
    labels = {
        "diamond": _("Diamond"),
        "green": _("Green"),
        "gold": _("Gold"),
        "hybrid": _("Hybrid"),
        "bronze": _("Bronze"),
        "closed": _("Closed"),
    }
    if not status:
        return _("Unknown OA status")
    return labels.get(status, status.replace("_", " ").replace("-", " ").title())


def _openalex_oa_status_color(value: str | None) -> str:
    """Return the IGI Global color associated with an Open Access status."""
    status = str(value or "").strip().lower()
    return _OPENALEX_OA_COLORS.get(status, "#a0a0a0")


def _openalex_language_label(value: str | None, locale=None) -> str:
    """Expand an OpenAlex ISO 639-1 language code in the active interface locale."""
    code = str(value or "").strip().lower()
    if not code:
        return _("Unknown language")

    display_locale = locale or get_locale()
    try:
        parsed_locale = (
            display_locale
            if isinstance(display_locale, Locale)
            else Locale.parse(
                str(display_locale or "en"),
                sep="_" if "_" in str(display_locale or "") else "-",
            )
        )
        language_name = parsed_locale.languages.get(code)
    except (UnknownLocaleError, ValueError):
        language_name = None

    if not language_name:
        return code.upper()
    return f"{language_name[:1].upper()}{language_name[1:]} ({code})"


def _empty_priority_open_access() -> dict:
    pagination = _pagination_dict(1, 10, 0)
    return {
        "summary": {
            "diamond_articles": 0,
            "diamond_citations": 0,
            "green_articles": 0,
            "green_citations": 0,
            "articles": 0,
            "citations": 0,
        },
        "source_rows": [],
        "top_cited_articles": [],
        "tables": {
            "sources": {
                "pagination": pagination,
                "per_page_options": [10, 25, 50],
                "search": "",
                "sort": "articles",
                "dir": "desc",
            },
            "articles": {
                "pagination": pagination,
                "per_page_options": [10, 25, 50],
                "search": "",
                "sort": "citations",
                "dir": "desc",
            },
            "institutions": {
                "pagination": pagination,
                "per_page_options": [10, 25, 50],
                "search": "",
                "sort": "articles",
                "dir": "desc",
            },
        },
        "charts": {
            "sources_by_articles": {"labels": [], "datasets": []},
            "sources_by_citations": {"labels": [], "datasets": []},
            "institutions_by_articles": {"labels": [], "datasets": []},
            "institutions_by_citations": {"labels": [], "datasets": []},
            "trend_years": [],
            "article_trend_datasets": [],
            "citation_trend_datasets": [],
        },
    }


def _priority_open_access_breakdown(doi_subquery) -> dict:
    """Aggregate Diamond and Green articles, sources, and citations."""
    from ..models import OpenAlexWorkMetadata

    source_page, source_per_page = _table_page_params(
        "priority_source",
        default_per_page=10,
        max_per_page=50,
    )
    article_page, article_per_page = _table_page_params(
        "priority_article",
        default_per_page=10,
        max_per_page=50,
    )
    source_sort, source_dir = _table_sort_params(
        "priority_source",
        {
            "source",
            "oa_status",
            "articles",
            "citations",
            "average_citations",
        },
        "articles",
    )
    article_sort, article_dir = _table_sort_params(
        "priority_article",
        {
            "article",
            "oa_status",
            "citations",
            "year",
            "openalex",
        },
        "citations",
    )
    source_search = (
        request.args.get("priority_source_q") or ""
    ).strip()
    article_search = (
        request.args.get("priority_article_q") or ""
    ).strip()
    priority_statuses = ("diamond", "green")
    status_expr = func.lower(OpenAlexWorkMetadata.oa_status)
    source_name_expr = func.coalesce(
        func.nullif(func.trim(OpenAlexWorkMetadata.source_name), ""),
        _("Unknown source"),
    )
    article_count_expr = func.count(func.distinct(OpenAlexWorkMetadata.doi_normalized))
    citation_count_expr = func.coalesce(func.sum(OpenAlexWorkMetadata.cited_by_count), 0)

    summary_rows = (
        db.session.query(
            status_expr.label("oa_status"),
            article_count_expr.label("article_count"),
            citation_count_expr.label("citation_count"),
        )
        .join(doi_subquery, OpenAlexWorkMetadata.doi_normalized == doi_subquery.c.doi_normalized)
        .filter(status_expr.in_(priority_statuses))
        .group_by(status_expr)
        .all()
    )
    summary_by_status = {
        row.oa_status: {
            "articles": int(row.article_count or 0),
            "citations": int(row.citation_count or 0),
        }
        for row in summary_rows
    }

    grouped_source_rows = (
        db.session.query(
            source_name_expr.label("source_name"),
            OpenAlexWorkMetadata.source_issn_l.label("source_issn_l"),
            status_expr.label("oa_status"),
            article_count_expr.label("article_count"),
            citation_count_expr.label("citation_count"),
        )
        .join(doi_subquery, OpenAlexWorkMetadata.doi_normalized == doi_subquery.c.doi_normalized)
        .filter(status_expr.in_(priority_statuses))
        .group_by(
            source_name_expr,
            OpenAlexWorkMetadata.source_issn_l,
            status_expr,
        )
        .all()
    )

    source_rows = []
    source_totals = {}
    source_status_values = {}
    for row in grouped_source_rows:
        source_key = (row.source_name, row.source_issn_l or "")
        articles = int(row.article_count or 0)
        citations = int(row.citation_count or 0)
        source_rows.append({
            "source_name": row.source_name,
            "source_issn_l": row.source_issn_l,
            "oa_status": row.oa_status,
            "oa_status_label": _openalex_oa_status_label(row.oa_status),
            "articles": articles,
            "citations": citations,
            "average_citations": round(citations / articles, 1) if articles else 0,
        })
        totals = source_totals.setdefault(
            source_key,
            {"articles": 0, "citations": 0},
        )
        totals["articles"] += articles
        totals["citations"] += citations
        source_status_values[(source_key, row.oa_status)] = {
            "articles": articles,
            "citations": citations,
        }

    if source_search:
        source_search_value = source_search.casefold()
        source_rows = [
            row
            for row in source_rows
            if source_search_value in " ".join((
                row["source_name"],
                row["source_issn_l"] or "",
                row["oa_status"],
                row["oa_status_label"],
            )).casefold()
        ]

    source_rows.sort(
        key=lambda row: (
            row["source_name"].casefold(),
            row["source_issn_l"] or "",
            row["oa_status"],
        )
    )
    if source_sort == "source":
        source_rows.sort(
            key=lambda row: (
                row["source_name"].casefold(),
                row["source_issn_l"] or "",
            ),
            reverse=source_dir == "desc",
        )
    elif source_sort == "oa_status":
        source_rows.sort(
            key=lambda row: row["oa_status_label"].casefold(),
            reverse=source_dir == "desc",
        )
    else:
        source_rows.sort(
            key=lambda row: row[source_sort],
            reverse=source_dir == "desc",
        )
    source_pagination = _pagination_dict(
        source_page,
        source_per_page,
        len(source_rows),
    )
    source_start = (
        source_pagination["page"] - 1
    ) * source_pagination["per_page"]
    paged_source_rows = source_rows[
        source_start:source_start + source_pagination["per_page"]
    ]

    def source_chart(metric: str) -> dict:
        ranked_keys = sorted(
            source_totals,
            key=lambda key: (
                -source_totals[key][metric],
                key[0].lower(),
                key[1],
            ),
        )[:10]
        return {
            "labels": [key[0] for key in ranked_keys],
            "datasets": [
                {
                    "label": _openalex_oa_status_label(status),
                    "data": [
                        source_status_values.get((key, status), {}).get(metric, 0)
                        for key in ranked_keys
                    ],
                    "backgroundColor": color,
                }
                for status, color in _OPENALEX_PRIORITY_OA_COLORS.items()
            ],
        }

    priority_article_query = (
        OpenAlexWorkMetadata.query
        .join(doi_subquery, OpenAlexWorkMetadata.doi_normalized == doi_subquery.c.doi_normalized)
        .filter(status_expr.in_(priority_statuses))
    )
    if article_search:
        article_search_value = article_search.casefold()
        article_search_like = f"%{article_search_value}%"
        matching_statuses = [
            status
            for status in priority_statuses
            if article_search_value in str(
                _openalex_oa_status_label(status)
            ).casefold()
        ]
        article_search_conditions = [
            func.lower(func.coalesce(OpenAlexWorkMetadata.title, "")).like(
                article_search_like
            ),
            func.lower(func.coalesce(OpenAlexWorkMetadata.source_name, "")).like(
                article_search_like
            ),
            func.lower(OpenAlexWorkMetadata.doi_normalized).like(
                article_search_like
            ),
            func.lower(func.coalesce(OpenAlexWorkMetadata.openalex_id, "")).like(
                article_search_like
            ),
        ]
        if matching_statuses:
            article_search_conditions.append(
                status_expr.in_(matching_statuses)
            )
        priority_article_query = priority_article_query.filter(
            or_(*article_search_conditions)
        )

    article_pagination = _pagination_dict(
        article_page,
        article_per_page,
        priority_article_query.count(),
    )
    article_sort_columns = {
        "article": func.lower(func.coalesce(OpenAlexWorkMetadata.title, "")),
        "oa_status": status_expr,
        "citations": OpenAlexWorkMetadata.cited_by_count,
        "year": func.coalesce(OpenAlexWorkMetadata.publication_year, 0),
        "openalex": func.lower(func.coalesce(OpenAlexWorkMetadata.openalex_id, "")),
    }
    article_order = article_sort_columns[article_sort]
    article_order = (
        article_order.asc()
        if article_dir == "asc"
        else article_order.desc()
    )
    top_cited_rows = (
        priority_article_query
        .order_by(
            article_order,
            OpenAlexWorkMetadata.cited_by_count.desc(),
            OpenAlexWorkMetadata.title.asc(),
        )
        .offset(
            (article_pagination["page"] - 1)
            * article_pagination["per_page"]
        )
        .limit(article_pagination["per_page"])
        .all()
    )
    top_cited_articles = [
        {
            "title": row.title,
            "doi": row.doi_normalized,
            "openalex_id": row.openalex_id,
            "publication_year": row.publication_year,
            "source_name": row.source_name,
            "source_issn_l": row.source_issn_l,
            "oa_status": str(row.oa_status or "").lower(),
            "oa_status_label": _openalex_oa_status_label(row.oa_status),
            "citations": int(row.cited_by_count or 0),
            "fwci": row.fwci,
        }
        for row in top_cited_rows
    ]

    trend_rows = (
        db.session.query(
            OpenAlexWorkMetadata.publication_year.label("publication_year"),
            status_expr.label("oa_status"),
            article_count_expr.label("article_count"),
            citation_count_expr.label("citation_count"),
        )
        .join(doi_subquery, OpenAlexWorkMetadata.doi_normalized == doi_subquery.c.doi_normalized)
        .filter(
            OpenAlexWorkMetadata.publication_year.isnot(None),
            status_expr.in_(priority_statuses),
        )
        .group_by(OpenAlexWorkMetadata.publication_year, status_expr)
        .order_by(OpenAlexWorkMetadata.publication_year.asc())
        .all()
    )
    trend_years = sorted({
        str(row.publication_year)
        for row in trend_rows
    })
    trend_values = {
        (str(row.publication_year), row.oa_status): {
            "articles": int(row.article_count or 0),
            "citations": int(row.citation_count or 0),
        }
        for row in trend_rows
    }

    def trend_datasets(metric: str) -> list[dict]:
        return [
            {
                "label": _openalex_oa_status_label(status),
                "data": [
                    trend_values.get((year, status), {}).get(metric, 0)
                    for year in trend_years
                ],
                "borderColor": color,
                "backgroundColor": color,
                "tension": 0.25,
                "fill": False,
            }
            for status, color in _OPENALEX_PRIORITY_OA_COLORS.items()
        ]

    diamond = summary_by_status.get("diamond", {"articles": 0, "citations": 0})
    green = summary_by_status.get("green", {"articles": 0, "citations": 0})
    return {
        "summary": {
            "diamond_articles": diamond["articles"],
            "diamond_citations": diamond["citations"],
            "green_articles": green["articles"],
            "green_citations": green["citations"],
            "articles": diamond["articles"] + green["articles"],
            "citations": diamond["citations"] + green["citations"],
        },
        "source_rows": paged_source_rows,
        "top_cited_articles": top_cited_articles,
        "tables": {
            "sources": {
                "pagination": source_pagination,
                "per_page_options": [10, 25, 50],
                "search": source_search,
                "sort": source_sort,
                "dir": source_dir,
            },
            "articles": {
                "pagination": article_pagination,
                "per_page_options": [10, 25, 50],
                "search": article_search,
                "sort": article_sort,
                "dir": article_dir,
            },
        },
        "charts": {
            "sources_by_articles": source_chart("articles"),
            "sources_by_citations": source_chart("citations"),
            "trend_years": trend_years,
            "article_trend_datasets": trend_datasets("articles"),
            "citation_trend_datasets": trend_datasets("citations"),
        },
    }
