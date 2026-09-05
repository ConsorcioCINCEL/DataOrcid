"""Institution-scoped OAI-PMH publication and provider routes."""

import csv
from io import BytesIO, StringIO
import re
from zipfile import BadZipFile

from flask import (
    abort,
    Blueprint,
    Response,
    current_app,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from flask_babel import _, force_locale, get_locale
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import and_, func, or_
from werkzeug.datastructures import CombinedMultiDict
from werkzeug.utils import secure_filename

from .. import DEFAULT_LANGUAGES, csrf, db
from ..decorators import institution_required, oai_editor_required, staff_required
from ..models import (
    CanonicalWork,
    OaiPmhDoiImportBatch,
    OaiPmhDoiImportChange,
    OaiPmhInstitutionConfig,
    OaiPmhHarvester,
    OaiPmhWorkSelection,
    OpenAlexInstitutionWorkFact,
    OpenAlexWorkMetadata,
    ResearcherCache,
    User,
    WorkCache,
    WorkRecordLink,
    utc_now,
)
from ..services.institution_registry_service import get_institution_by_ror
from ..services.doi_service import normalize_doi
from ..services.oai_access import MAX_HARVESTERS, normalize_harvester_uri
from ..services.oai_pmh_service import (
    METADATA_FIELD_CATALOG,
    METADATA_SOURCE_FIELDS,
    MAPPING_TARGET_RE,
    attach_institutional_affiliations,
    build_oai_response,
    effective_metadata_mapping,
    generate_public_key,
    institutional_work_query,
    work_from_query_row,
)
from ..utils.flashes import flash_err, flash_ok, flash_warn


bp_oai_pmh = Blueprint("oai_pmh", __name__)
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+$")
PUBLICATION_POLICIES = {"validated", "all", "selected"}
DOI_IMPORT_MAX_BYTES = 5 * 1024 * 1024
DOI_IMPORT_MAX_ROWS = 5000
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXCEL_ILLEGAL_CHARACTERS_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")


class DoiImportError(ValueError):
    """Represent a user-correctable DOI spreadsheet validation error."""


@bp_oai_pmh.route("/oai-pmh/")
@institution_required
def index():
    """Render the initial institutional OAI-PMH configuration workspace."""
    ror_id = g.institution_ror_id
    institution = get_institution_by_ror(ror_id) or {}
    institution_name = institution.get("name") or session.get("institution_name") or ror_id
    stored_config = OaiPmhInstitutionConfig.query.filter_by(ror_id=ror_id).first()
    config = stored_config or OaiPmhInstitutionConfig(
        ror_id=ror_id,
        repository_name=f"DataORCID-Chile — {institution_name}",
        public_key="",
        provider_enabled=False,
        publication_policy="validated",
        policy_updated_at=utc_now().replace(microsecond=0),
        created_at=utc_now().replace(microsecond=0),
        updated_at=utc_now().replace(microsecond=0),
    )
    count_query, count_included_expr, _count_datestamp = institutional_work_query(config)
    total = count_query.count()
    included = count_query.filter(count_included_expr.is_(True)).count()
    validated = count_query.filter(
        OpenAlexInstitutionWorkFact.has_selected_affiliation.is_(True)
    ).count()
    explicit = OaiPmhWorkSelection.query.filter_by(ror_id=ror_id).count()
    stats = {
        "total": total,
        "included": included,
        "excluded": max(total - included, 0),
        "validated": validated,
        "unvalidated": max(total - validated, 0),
        "explicit": explicit,
    }
    provider_url = _provider_url(stored_config.public_key) if stored_config else None

    return render_template(
        "oai_pmh/index.html",
        oai_config=stored_config,
        settings=config,
        institution={"ror_id": ror_id, "name": institution_name},
        stats=stats,
        provider_url=provider_url,
        identify_url=f"{provider_url}?verb=Identify" if provider_url else None,
        records_url=f"{provider_url}?verb=ListRecords&metadataPrefix=oai_dc" if provider_url else None,
        openaire_url=f"{provider_url}?verb=ListRecords&metadataPrefix=oai_openaire" if provider_url else None,
        mapped_url=f"{provider_url}?verb=ListRecords&metadataPrefix=dataorcid" if provider_url else None,
        active_tab="configuration",
        can_manage_provider=_can_manage_provider(),
        can_manage_content=_can_manage_oai_content(),
    )


@bp_oai_pmh.route("/oai-pmh/articles/")
@institution_required
def articles():
    """Render the filterable institutional OAI article-selection workspace."""
    ror_id = g.institution_ror_id
    institution = get_institution_by_ror(ror_id) or {}
    institution_name = institution.get("name") or session.get("institution_name") or ror_id
    stored_config = OaiPmhInstitutionConfig.query.filter_by(ror_id=ror_id).first()
    config = stored_config or OaiPmhInstitutionConfig(
        ror_id=ror_id,
        repository_name=f"DataORCID-Chile — {institution_name}",
        public_key="",
        provider_enabled=False,
        publication_policy="validated",
        policy_updated_at=utc_now().replace(microsecond=0),
        created_at=utc_now().replace(microsecond=0),
        updated_at=utc_now().replace(microsecond=0),
    )
    listing = _article_listing(config, request.args)
    works_query = listing["query"]
    page = max(_safe_int(request.args.get("page"), 1), 1)
    pagination = works_query.order_by(
        listing["ordering"],
        CanonicalWork.id.desc(),
    ).paginate(page=page, per_page=25, error_out=False)
    records = [work_from_query_row(row, config) for row in pagination.items]
    attach_institutional_affiliations(records, ror_id)

    def article_sort_url(column: str) -> str:
        direction = (
            "desc"
            if listing["sort"] == column and listing["direction"] == "asc"
            else "asc"
        )
        return url_for(
            "oai_pmh.articles",
            q=listing["query_text"],
            status=listing["status"],
            type=listing["type"],
            validation=listing["validation"],
            source=listing["source"],
            sort=column,
            direction=direction,
        )

    return render_template(
        "oai_pmh/articles.html",
        institution={"ror_id": ror_id, "name": institution_name},
        records=records,
        pagination=pagination,
        query=listing["query_text"],
        selected_status=listing["status"],
        selected_type=listing["type"],
        selected_validation=listing["validation"],
        selected_source=listing["source"],
        selected_sort=listing["sort"],
        selected_direction=listing["direction"],
        article_sort_url=article_sort_url,
        type_options=listing["type_options"],
        active_tab="articles",
        can_manage_content=_can_manage_oai_content(),
    )


@bp_oai_pmh.route("/oai-pmh/metadata/")
@institution_required
def metadata_mapping():
    """Render the separate metadata-format and custom crosswalk workspace."""
    ror_id = g.institution_ror_id
    institution = get_institution_by_ror(ror_id) or {}
    institution_name = institution.get("name") or session.get("institution_name") or ror_id
    stored_config = OaiPmhInstitutionConfig.query.filter_by(ror_id=ror_id).first()
    config = stored_config or OaiPmhInstitutionConfig(
        ror_id=ror_id,
        repository_name=f"DataORCID-Chile — {institution_name}",
        public_key="",
        provider_enabled=False,
        publication_policy="validated",
        policy_updated_at=utc_now().replace(microsecond=0),
        created_at=utc_now().replace(microsecond=0),
        updated_at=utc_now().replace(microsecond=0),
    )
    mapping = effective_metadata_mapping(config)
    catalog = _metadata_catalog_view(mapping)
    provider_url = _provider_url(stored_config.public_key) if stored_config else None
    format_urls = {
        "oai_dc": f"{provider_url}?verb=ListRecords&metadataPrefix=oai_dc"
        if stored_config and stored_config.provider_enabled and not stored_config.harvester_access_restricted
        else None,
        "oai_openaire": f"{provider_url}?verb=ListRecords&metadataPrefix=oai_openaire"
        if stored_config and stored_config.provider_enabled and not stored_config.harvester_access_restricted
        else None,
        "dataorcid": f"{provider_url}?verb=ListRecords&metadataPrefix=dataorcid"
        if stored_config and stored_config.provider_enabled and not stored_config.harvester_access_restricted
        else None,
    }
    return render_template(
        "oai_pmh/metadata_mapping.html",
        oai_config=stored_config,
        settings=config,
        institution={"ror_id": ror_id, "name": institution_name},
        format_urls=format_urls,
        metadata_catalog=catalog,
        metadata_group_labels=_metadata_group_labels(),
        active_metadata_count=sum(1 for field in catalog if field["active"]),
        active_tab="metadata",
        can_manage=_can_manage_oai_content(),
        can_manage_content=_can_manage_oai_content(),
    )


@bp_oai_pmh.route("/oai-pmh/doi-import/")
@institution_required
def doi_import_workspace():
    """Render DOI upload controls and the auditable institutional history."""
    ror_id = g.institution_ror_id
    institution = get_institution_by_ror(ror_id) or {}
    institution_name = institution.get("name") or session.get("institution_name") or ror_id
    page = max(_safe_int(request.args.get("page"), 1), 1)
    history = (
        OaiPmhDoiImportBatch.query
        .filter_by(ror_id=ror_id)
        .order_by(
            OaiPmhDoiImportBatch.created_at.desc(),
            OaiPmhDoiImportBatch.id.desc(),
        )
        .paginate(page=page, per_page=20, error_out=False)
    )
    latest_active = (
        OaiPmhDoiImportBatch.query
        .filter_by(ror_id=ror_id, undone_at=None)
        .order_by(
            OaiPmhDoiImportBatch.created_at.desc(),
            OaiPmhDoiImportBatch.id.desc(),
        )
        .first()
    )
    actor_ids = {
        actor_id
        for batch in history.items
        for actor_id in (batch.imported_by_user_id, batch.undone_by_user_id)
        if actor_id
    }
    return render_template(
        "oai_pmh/doi_import.html",
        institution={"ror_id": ror_id, "name": institution_name},
        history=history,
        actor_names=_user_display_names(actor_ids),
        latest_active_id=latest_active.id if latest_active else None,
        active_tab="doi_import",
        can_manage_content=_can_manage_oai_content(),
    )


@bp_oai_pmh.route("/oai-pmh/doi-import/<int:batch_id>/")
@institution_required
def doi_import_detail(batch_id: int):
    """Show every article-level change recorded for one DOI import batch."""
    batch = OaiPmhDoiImportBatch.query.filter_by(
        id=batch_id,
        ror_id=g.institution_ror_id,
    ).first()
    if not batch:
        abort(404)
    page = max(_safe_int(request.args.get("page"), 1), 1)
    changes = (
        db.session.query(
            OaiPmhDoiImportChange,
            CanonicalWork,
            OaiPmhWorkSelection,
        )
        .join(
            CanonicalWork,
            CanonicalWork.id == OaiPmhDoiImportChange.canonical_work_id,
        )
        .outerjoin(
            OaiPmhWorkSelection,
            and_(
                OaiPmhWorkSelection.ror_id == g.institution_ror_id,
                OaiPmhWorkSelection.canonical_work_id == CanonicalWork.id,
            ),
        )
        .filter(OaiPmhDoiImportChange.batch_id == batch.id)
        .order_by(CanonicalWork.title.asc().nullslast(), CanonicalWork.id.asc())
        .paginate(page=page, per_page=50, error_out=False)
    )
    latest_active = (
        OaiPmhDoiImportBatch.query
        .filter_by(ror_id=g.institution_ror_id, undone_at=None)
        .order_by(
            OaiPmhDoiImportBatch.created_at.desc(),
            OaiPmhDoiImportBatch.id.desc(),
        )
        .first()
    )
    actor_ids = {
        actor_id
        for actor_id in (batch.imported_by_user_id, batch.undone_by_user_id)
        if actor_id
    }
    institution = get_institution_by_ror(g.institution_ror_id) or {}
    return render_template(
        "oai_pmh/doi_import_detail.html",
        institution={
            "ror_id": g.institution_ror_id,
            "name": institution.get("name")
            or session.get("institution_name")
            or g.institution_ror_id,
        },
        batch=batch,
        changes=changes,
        actor_names=_user_display_names(actor_ids),
        can_undo=(
            _can_manage_oai_content()
            and batch.undone_at is None
            and latest_active is not None
            and latest_active.id == batch.id
        ),
        active_tab="doi_import",
        can_manage_content=_can_manage_oai_content(),
    )


@bp_oai_pmh.route("/oai-pmh/access/")
@institution_required
def harvesting_access():
    """Show repositories belonging to the authenticated institution."""
    ror_id = g.institution_ror_id
    institution = get_institution_by_ror(ror_id) or {}
    config = OaiPmhInstitutionConfig.query.filter_by(ror_id=ror_id).first()
    harvesters = OaiPmhHarvester.query.filter_by(config_id=config.id).order_by(
        OaiPmhHarvester.created_at, OaiPmhHarvester.id,
    ).all() if config else []
    can_manage = _can_manage_oai_content()
    response = current_app.make_response(render_template(
        "oai_pmh/harvesting_access.html",
        institution={"ror_id": ror_id, "name": institution.get("name") or ror_id},
        oai_config=config,
        harvesters=harvesters,
        harvester_urls={item.id: _provider_url(config.public_key, item.access_key)
                        for item in harvesters if item.is_enabled} if can_manage else {},
        can_manage_content=can_manage,
        active_tab="access",
    ))
    response.headers.update({"Cache-Control": "private, no-store", "Referrer-Policy": "no-referrer"})
    return response


@bp_oai_pmh.route("/oai-pmh/access/", methods=["POST"])
@oai_editor_required
@institution_required
def save_harvesting_access():
    """Register a private URL or change only this institution's access mode."""
    # Serialize changes within one institution, including the repository limit.
    config = OaiPmhInstitutionConfig.query.filter_by(
        ror_id=g.institution_ror_id,
    ).with_for_update().first()
    if config is None:
        flash_err(_("Ask the responsible team to configure this OAI-PMH provider before managing harvesting access."))
        return redirect(url_for("oai_pmh.harvesting_access"))
    action = request.form.get("action")
    try:
        if action == "register":
            uri = normalize_harvester_uri(request.form.get("base_uri", ""))
            if OaiPmhHarvester.query.filter_by(config_id=config.id, base_uri=uri).first():
                raise ValueError(_("Each repository URI must appear only once."))
            if OaiPmhHarvester.query.filter_by(config_id=config.id).count() >= MAX_HARVESTERS:
                raise ValueError(_("You can register up to 20 repositories per institution."))
            db.session.add(OaiPmhHarvester(
                config_id=config.id, base_uri=uri,
                created_by_user_id=session.get("user_id"),
                updated_by_user_id=session.get("user_id"),
            ))
        elif action == "policy":
            restricted = request.form.get("restrict_access") == "on"
            if restricted and not OaiPmhHarvester.query.filter_by(config_id=config.id, is_enabled=True).first():
                raise ValueError(_("Add at least one active repository before restricting harvesting access."))
            config.harvester_access_restricted = restricted
        else:
            abort(400)
    except ValueError as exc:
        db.session.rollback()
        flash_err(str(exc))
        return redirect(url_for("oai_pmh.harvesting_access"))
    config.updated_by_user_id = session.get("user_id")
    config.updated_at = utc_now().replace(microsecond=0)
    db.session.commit()
    flash_ok(_("Harvesting access updated for this institution."))
    return redirect(url_for("oai_pmh.harvesting_access"))


@bp_oai_pmh.route("/oai-pmh/access/<int:harvester_id>/", methods=["POST"])
@oai_editor_required
@institution_required
def update_harvester(harvester_id: int):
    """Revoke, replace, or remove a credential without crossing institution scope."""
    config = OaiPmhInstitutionConfig.query.filter_by(
        ror_id=g.institution_ror_id,
    ).with_for_update().first_or_404()
    harvester = OaiPmhHarvester.query.filter_by(id=harvester_id, config_id=config.id).first_or_404()
    action = request.form.get("action")
    if action == "revoke":
        harvester.is_enabled = False
    elif action == "rotate":
        harvester.access_key = generate_public_key()
        harvester.is_enabled = True
    elif action == "delete":
        db.session.delete(harvester)
    else:
        abort(400)
    # Revoking the final credential keeps restricted mode closed.
    harvester.updated_by_user_id = session.get("user_id")
    harvester.updated_at = utc_now().replace(microsecond=0)
    config.updated_by_user_id = session.get("user_id")
    config.updated_at = harvester.updated_at
    db.session.commit()
    flash_ok(_("Harvesting access updated for this institution."))
    return redirect(url_for("oai_pmh.harvesting_access"))


@bp_oai_pmh.route("/oai-pmh/settings", methods=["POST"])
@staff_required
@institution_required
def save_settings():
    """Save provider identity and the default publication policy for one ROR."""
    ror_id = g.institution_ror_id
    institution = get_institution_by_ror(ror_id) or {}
    config = OaiPmhInstitutionConfig.query.filter_by(ror_id=ror_id).first()
    repository_name = (request.form.get("repository_name") or "").strip()
    admin_email = (request.form.get("admin_email") or "").strip()
    publication_policy = (request.form.get("publication_policy") or "validated").strip()
    provider_enabled = request.form.get("provider_enabled") == "on"

    try:
        if not repository_name or len(repository_name) > 255:
            raise ValueError(_("Enter a repository name of 255 characters or fewer."))
        if admin_email and (len(admin_email) > 255 or not EMAIL_RE.fullmatch(admin_email)):
            raise ValueError(_("Enter a valid repository administrator email."))
        if provider_enabled and not admin_email:
            raise ValueError(_("An administrator email is required to enable the OAI-PMH provider."))
        if publication_policy not in PUBLICATION_POLICIES:
            raise ValueError(_("Choose a valid default publication policy."))
    except ValueError as exc:
        flash_err(str(exc))
        return redirect(url_for("oai_pmh.index"))

    now = utc_now().replace(microsecond=0)
    if not config:
        config = OaiPmhInstitutionConfig(
            ror_id=ror_id,
            public_key=generate_public_key(),
            repository_name=repository_name or institution.get("name") or ror_id,
            publication_policy=publication_policy,
            policy_updated_at=now,
            created_by_user_id=session.get("user_id"),
            created_at=now,
            updated_at=now,
        )
        db.session.add(config)
    elif config.publication_policy != publication_policy:
        config.policy_updated_at = now

    config.provider_enabled = provider_enabled
    config.repository_name = repository_name
    config.admin_email = admin_email or None
    config.publication_policy = publication_policy
    config.updated_by_user_id = session.get("user_id")
    config.updated_at = now
    try:
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        current_app.logger.exception("Could not save OAI-PMH settings for %s: %s", ror_id, exc)
        flash_err(_("Could not save the OAI-PMH configuration."))
    else:
        flash_ok(_("OAI-PMH provider configuration saved."))
    return redirect(url_for("oai_pmh.index"))


@bp_oai_pmh.route("/oai-pmh/access-key/rotate", methods=["POST"])
@staff_required
@institution_required
def rotate_access_key():
    """Replace the opaque public URL and immediately invalidate the old one."""
    config = OaiPmhInstitutionConfig.query.filter_by(ror_id=g.institution_ror_id).first()
    if not config:
        flash_err(_("Save the OAI-PMH provider configuration before rotating its public key."))
        return redirect(url_for("oai_pmh.index"))

    config.public_key = generate_public_key()
    now = utc_now().replace(microsecond=0)
    config.updated_by_user_id = session.get("user_id")
    config.policy_updated_at = now
    config.updated_at = now
    db.session.commit()
    flash_ok(_("The OAI-PMH public key was rotated. The previous URL no longer works."))
    return redirect(url_for("oai_pmh.index"))


@bp_oai_pmh.route("/oai-pmh/metadata-mapping", methods=["POST"])
@oai_editor_required
@institution_required
def save_metadata_mapping():
    """Save aliases used by the institution-specific DataORCID metadata format."""
    config = OaiPmhInstitutionConfig.query.filter_by(ror_id=g.institution_ror_id).first()
    if not config:
        flash_err(_("Save the OAI-PMH provider configuration before editing its metadata mapping."))
        return redirect(url_for("oai_pmh.metadata_mapping"))

    if request.form.get("mapping_action") == "reset":
        config.metadata_mapping = None
        success_message = _("The default metadata mapping was restored.")
    else:
        mapping = {}
        submitted_fields = [
            field
            for field in request.form.getlist("mapping_fields")
            if field in METADATA_SOURCE_FIELDS
        ]
        if not submitted_fields:
            # Preserve compatibility with mappings submitted by the previous
            # fixed-field form and with API clients built against it.
            submitted_fields = [
                field
                for field in METADATA_SOURCE_FIELDS
                if f"mapping_{field}" in request.form
            ]
        for field in dict.fromkeys(submitted_fields):
            target = (request.form.get(f"mapping_{field}") or "").strip()
            if target and not MAPPING_TARGET_RE.fullmatch(target):
                flash_err(_(
                    "The target name for %(field)s is invalid. Use letters, numbers, dots, "
                    "colons, underscores, or hyphens.",
                    field=field,
                ))
                return redirect(url_for("oai_pmh.metadata_mapping"))
            if target:
                mapping[field] = target
        if "title" not in mapping or "identifier" not in mapping:
            flash_err(_("The custom mapping must keep target names for title and identifier."))
            return redirect(url_for("oai_pmh.metadata_mapping"))
        config.metadata_mapping = mapping
        success_message = _("Institutional metadata mapping saved.")

    now = utc_now().replace(microsecond=0)
    config.updated_by_user_id = session.get("user_id")
    config.policy_updated_at = now
    config.updated_at = now
    db.session.commit()
    flash_ok(success_message)
    return redirect(url_for("oai_pmh.metadata_mapping"))


@bp_oai_pmh.route("/oai-pmh/works/publication", methods=["POST"])
@oai_editor_required
@institution_required
def update_publication():
    """Include or exclude selected DataORCID works inside the active ROR."""
    ror_id = g.institution_ror_id
    single_include = _safe_int(request.form.get("single_include"), 0)
    single_exclude = _safe_int(request.form.get("single_exclude"), 0)
    if single_include > 0:
        action = "include"
        work_ids = {single_include}
    elif single_exclude > 0:
        action = "exclude"
        work_ids = {single_exclude}
    else:
        action = (request.form.get("action") or "").strip()
        work_ids = {
            value
            for value in (_safe_int(item, 0) for item in request.form.getlist("work_ids"))
            if value > 0
        }
    if action not in {"include", "exclude"} or not work_ids:
        flash_err(_("Select at least one article and a valid publication action."))
        return redirect(_return_to_articles())

    allowed_ids = {
        value
        for (value,) in db.session.query(WorkRecordLink.canonical_work_id)
        .filter(
            WorkRecordLink.ror_id == ror_id,
            WorkRecordLink.canonical_work_id.in_(work_ids),
        )
        .distinct()
        .all()
    }
    now = utc_now().replace(microsecond=0)
    desired = action == "include"
    selections = {
        row.canonical_work_id: row
        for row in OaiPmhWorkSelection.query.filter(
            OaiPmhWorkSelection.ror_id == ror_id,
            OaiPmhWorkSelection.canonical_work_id.in_(allowed_ids),
        ).all()
    }
    for work_id in allowed_ids:
        selection = selections.get(work_id)
        if not selection:
            selection = OaiPmhWorkSelection(
                ror_id=ror_id,
                canonical_work_id=work_id,
                created_at=now,
            )
            db.session.add(selection)
        selection.is_included = desired
        selection.decision_source = "manual"
        selection.doi_import_batch_id = None
        selection.updated_by_user_id = session.get("user_id")
        selection.updated_at = now
    db.session.commit()

    if len(allowed_ids) != len(work_ids):
        flash_err(_("Some articles were ignored because they do not belong to this institution."))
    if desired:
        flash_ok(_("Included %(count)s articles in the OAI-PMH provider.", count=len(allowed_ids)))
    else:
        flash_ok(_("Excluded %(count)s articles from the OAI-PMH provider.", count=len(allowed_ids)))
    return redirect(_return_to_articles())


@bp_oai_pmh.route("/oai-pmh/works/doi-import/template")
@oai_editor_required
@institution_required
def download_doi_import_template():
    """Download a guided XLSX template for institutional DOI activation."""
    workbook = Workbook()
    doi_sheet = workbook.active
    doi_sheet.title = "DOI"
    doi_sheet.append(["doi", _("Optional note")])
    doi_sheet.freeze_panes = "A2"
    doi_sheet.auto_filter.ref = "A1:B1"
    doi_sheet.column_dimensions["A"].width = 46
    doi_sheet.column_dimensions["B"].width = 42

    header_fill = PatternFill("solid", fgColor="C94725")
    for cell in doi_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    doi_sheet["A1"].comment = Comment(
        _("Enter one DOI per row. DOI URLs and the doi: prefix are accepted."),
        "DataORCID-Chile",
    )

    instructions = workbook.create_sheet(_("Instructions"))
    instructions.column_dimensions["A"].width = 105
    instructions.append([_("OAI-PMH bulk activation by DOI")])
    instructions.append([_("How to complete this template")])
    instructions.append([_("1. Open the DOI sheet and enter one DOI per row under the doi column.")])
    instructions.append([_("2. Keep the doi header unchanged. The optional note column is ignored during import.")])
    instructions.append([_("3. Save the file as XLSX and upload it from the OAI-PMH bulk DOI activation page.")])
    instructions.append([_("Only articles already associated with the active institution can be activated.")])
    instructions.append([_("Accepted example: 10.1234/example or https://doi.org/10.1234/example")])
    instructions["A1"].font = Font(size=14, bold=True, color="C94725")
    instructions["A2"].font = Font(bold=True)
    instructions.freeze_panes = "A3"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="plantilla_activacion_oai_doi.xlsx",
        mimetype=XLSX_MIMETYPE,
    )


@bp_oai_pmh.route("/oai-pmh/works/export")
@institution_required
def export_article_listing():
    """Export every filtered institutional article for auditing."""
    export_format = (request.args.get("format") or "xlsx").strip().lower()
    if export_format not in {"csv", "xlsx"}:
        export_format = "xlsx"

    ror_id = g.institution_ror_id
    institution = get_institution_by_ror(ror_id) or {}
    institution_name = institution.get("name") or session.get("institution_name") or ror_id
    stored_config = OaiPmhInstitutionConfig.query.filter_by(ror_id=ror_id).first()
    config = stored_config or OaiPmhInstitutionConfig(
        ror_id=ror_id,
        repository_name=f"DataORCID-Chile — {institution_name}",
        public_key="",
        provider_enabled=False,
        publication_policy="validated",
        policy_updated_at=utc_now().replace(microsecond=0),
        created_at=utc_now().replace(microsecond=0),
        updated_at=utc_now().replace(microsecond=0),
    )
    if request.args.get("background") == "1":
        from ..services.export_jobs import queue_export_response

        return queue_export_response(
            "oai_article_audit",
            export_format,
            {
                "ror_id": ror_id,
                "institution_name": institution_name,
                "filters": {
                    key: request.args.get(key)
                    for key in ("q", "status", "type", "validation", "source", "sort", "direction")
                },
                "locale": session.get("locale") or current_app.config.get("BABEL_DEFAULT_LOCALE", "en"),
            },
            _("OAI article audit"),
        )
    listing = _article_listing(config, request.args, load_type_options=False)
    ordered_query = listing["query"].order_by(
        listing["ordering"],
        CanonicalWork.id.desc(),
    )
    headers = _article_export_headers()
    rows = _iter_article_export_rows(ordered_query, config, ror_id)
    timestamp = utc_now().strftime("%Y%m%d_%H%M%S")
    safe_ror = re.sub(r"[^A-Za-z0-9_-]", "_", ror_id)

    if export_format == "csv":
        text_output = StringIO(newline="")
        writer = csv.writer(text_output)
        writer.writerow(headers)
        exported_count = 0
        for row in rows:
            writer.writerow([_safe_spreadsheet_cell(value) for value in row])
            exported_count += 1
        output = BytesIO(text_output.getvalue().encode("utf-8-sig"))
        current_app.logger.info(
            "Exported %s OAI audit rows as CSV for %s",
            exported_count,
            ror_id,
        )
        return send_file(
            output,
            as_attachment=True,
            download_name=f"oai_articulos_{safe_ror}_{timestamp}.csv",
            mimetype="text/csv; charset=utf-8",
        )

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(_("Articles"))
    worksheet.freeze_panes = "A2"
    for column, width in zip(
        "ABCDEFGHIJKLMN",
        [48, 28, 45, 10, 18, 30, 60, 45, 38, 20, 16, 22, 22, 48],
    ):
        worksheet.column_dimensions[column].width = width
    header_fill = PatternFill("solid", fgColor="C94725")
    header_cells = []
    for value in headers:
        cell = WriteOnlyCell(worksheet, value=value)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
        header_cells.append(cell)
    worksheet.append(header_cells)
    exported_count = 0
    for row in rows:
        worksheet.append([
            _safe_spreadsheet_cell(value, xlsx=True) for value in row
        ])
        exported_count += 1
    worksheet.auto_filter.ref = f"A1:N{max(exported_count + 1, 1)}"

    summary = workbook.create_sheet(_("Audit summary"))
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 70
    summary_rows = [
        [_("Institution"), institution_name],
        ["ROR", ror_id],
        [_("Generated at"), utc_now().strftime("%Y-%m-%dT%H:%M:%SZ")],
        [_("Exported articles"), exported_count],
        [_("Search filter"), listing["query_text"] or _("None")],
        [_("OAI status filter"), listing["status"]],
        [_("OpenAlex validation filter"), listing["validation"]],
        [_("Activation-origin filter"), listing["source"]],
    ]
    for summary_row in summary_rows:
        summary.append([
            _safe_spreadsheet_cell(value, xlsx=True) for value in summary_row
        ])

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    current_app.logger.info(
        "Exported %s OAI audit rows as XLSX for %s",
        exported_count,
        ror_id,
    )
    return send_file(
        output,
        as_attachment=True,
        download_name=f"oai_articulos_{safe_ror}_{timestamp}.xlsx",
        mimetype=XLSX_MIMETYPE,
    )


@bp_oai_pmh.route("/oai-pmh/works/doi-import", methods=["POST"])
@oai_editor_required
@institution_required
def import_doi_publication():
    """Activate institutional OAI articles listed by DOI in an XLSX file."""
    upload = request.files.get("doi_file")
    if not upload or not upload.filename:
        flash_err(_("Choose an XLSX file containing the DOI list."))
        return redirect(_return_to_doi_import())
    if not upload.filename.lower().endswith(".xlsx"):
        flash_err(_("The DOI import file must use the XLSX format."))
        return redirect(_return_to_doi_import())

    try:
        import_rows = _read_doi_import(upload)
    except DoiImportError as exc:
        flash_err(str(exc))
        return redirect(_return_to_doi_import())
    except (BadZipFile, InvalidFileException, EOFError, OSError, KeyError) as exc:
        current_app.logger.info("Rejected invalid OAI DOI spreadsheet: %s", exc)
        flash_err(_("The uploaded XLSX file could not be read. Download the template and try again."))
        return redirect(_return_to_doi_import())
    except Exception as exc:
        current_app.logger.exception("Could not parse the OAI DOI spreadsheet: %s", exc)
        flash_err(_("The DOI spreadsheet could not be processed."))
        return redirect(_return_to_doi_import())

    unique_dois = list(dict.fromkeys(import_rows["normalized_rows"]))
    matched_rows = (
        db.session.query(CanonicalWork.id, CanonicalWork.doi_normalized)
        .join(
            WorkRecordLink,
            WorkRecordLink.canonical_work_id == CanonicalWork.id,
        )
        .join(WorkCache, WorkCache.id == WorkRecordLink.work_cache_id)
        .filter(
            WorkRecordLink.ror_id == g.institution_ror_id,
            WorkCache.visibility == "public",
            CanonicalWork.doi_normalized.in_(unique_dois),
        )
        .distinct()
        .all()
    )
    matched_dois = {row.doi_normalized for row in matched_rows}
    allowed_ids = {row.id for row in matched_rows}
    matched_count = len(matched_dois)
    unmatched_count = len(unique_dois) - matched_count
    now = utc_now().replace(microsecond=0)
    selections = {
        row.canonical_work_id: row
        for row in OaiPmhWorkSelection.query.filter(
            OaiPmhWorkSelection.ror_id == g.institution_ror_id,
            OaiPmhWorkSelection.canonical_work_id.in_(allowed_ids),
        ).all()
    } if allowed_ids else {}

    filename = secure_filename(upload.filename)[:255] or "doi_import.xlsx"
    batch = OaiPmhDoiImportBatch(
        ror_id=g.institution_ror_id,
        filename=filename,
        submitted_count=import_rows["submitted_count"],
        matched_count=matched_count,
        article_count=len(allowed_ids),
        invalid_count=import_rows["invalid_count"],
        duplicate_count=import_rows["duplicate_count"],
        unmatched_count=unmatched_count,
        imported_by_user_id=session.get("user_id"),
        created_at=now,
    )
    db.session.add(batch)
    db.session.flush()

    if allowed_ids:
        for work_id in allowed_ids:
            selection = selections.get(work_id)
            db.session.add(OaiPmhDoiImportChange(
                batch_id=batch.id,
                canonical_work_id=work_id,
                previous_selection_existed=selection is not None,
                previous_is_included=(selection.is_included if selection else None),
                previous_decision_source=(selection.decision_source if selection else None),
                previous_import_batch_id=(selection.doi_import_batch_id if selection else None),
            ))
            if not selection:
                selection = OaiPmhWorkSelection(
                    ror_id=g.institution_ror_id,
                    canonical_work_id=work_id,
                    created_at=now,
                )
                db.session.add(selection)
            selection.is_included = True
            selection.decision_source = "xlsx"
            selection.doi_import_batch_id = batch.id
            selection.updated_by_user_id = session.get("user_id")
            selection.updated_at = now

    try:
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        current_app.logger.exception(
            "Could not import OAI DOI selections for %s: %s",
            g.institution_ror_id,
            exc,
        )
        flash_err(_("The DOI selections could not be saved."))
        return redirect(_return_to_doi_import())

    flash_ok(_(
        "Excel import complete: %(matched)s of %(total)s DOI rows matched this institution; "
        "%(articles)s articles were activated.",
        matched=matched_count,
        total=import_rows["submitted_count"],
        articles=len(allowed_ids),
    ))
    if import_rows["invalid_count"] or import_rows["duplicate_count"] or unmatched_count:
        flash_warn(_(
            "Ignored rows: %(invalid)s invalid DOI, %(duplicates)s duplicates, and "
            "%(unmatched)s without a matching institutional article.",
            invalid=import_rows["invalid_count"],
            duplicates=import_rows["duplicate_count"],
            unmatched=unmatched_count,
        ))
    return redirect(url_for("oai_pmh.doi_import_workspace"))


@bp_oai_pmh.route(
    "/oai-pmh/works/doi-import/undo",
    defaults={"batch_id": None},
    methods=["POST"],
)
@bp_oai_pmh.route(
    "/oai-pmh/works/doi-import/<int:batch_id>/undo",
    methods=["POST"],
)
@oai_editor_required
@institution_required
def undo_last_doi_import(batch_id: int | None):
    """Restore selections changed by the latest reversible DOI import batch."""
    batch = (
        OaiPmhDoiImportBatch.query
        .filter_by(ror_id=g.institution_ror_id, undone_at=None)
        .order_by(
            OaiPmhDoiImportBatch.created_at.desc(),
            OaiPmhDoiImportBatch.id.desc(),
        )
        .first()
    )
    if not batch or (batch_id is not None and batch.id != batch_id):
        flash_err(_("There is no DOI import available to undo."))
        return redirect(_return_to_doi_import())

    changes = OaiPmhDoiImportChange.query.filter_by(batch_id=batch.id).all()
    work_ids = {change.canonical_work_id for change in changes}
    selections = {
        row.canonical_work_id: row
        for row in OaiPmhWorkSelection.query.filter(
            OaiPmhWorkSelection.ror_id == g.institution_ror_id,
            OaiPmhWorkSelection.canonical_work_id.in_(work_ids),
        ).all()
    } if work_ids else {}
    now = utc_now().replace(microsecond=0)
    restored_count = 0
    skipped_count = 0
    for change in changes:
        selection = selections.get(change.canonical_work_id)
        if not selection or selection.doi_import_batch_id != batch.id:
            skipped_count += 1
            continue
        if change.previous_selection_existed:
            selection.is_included = bool(change.previous_is_included)
            selection.decision_source = change.previous_decision_source or "manual"
            selection.doi_import_batch_id = change.previous_import_batch_id
            selection.updated_by_user_id = session.get("user_id")
            selection.updated_at = now
        else:
            db.session.delete(selection)
        restored_count += 1

    batch.undone_at = now
    batch.undone_by_user_id = session.get("user_id")
    try:
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        current_app.logger.exception(
            "Could not undo OAI DOI import batch %s for %s: %s",
            batch.id,
            g.institution_ror_id,
            exc,
        )
        flash_err(_("The last DOI import could not be undone."))
        return redirect(_return_to_doi_import())

    flash_ok(_(
        "The last DOI import (%(filename)s) was undone for %(count)s articles.",
        filename=batch.filename,
        count=restored_count,
    ))
    if skipped_count:
        flash_warn(_(
            "%(count)s articles were not changed because they were edited after the import.",
            count=skipped_count,
        ))
    return redirect(url_for("oai_pmh.doi_import_workspace"))


def _article_listing(
    config,
    values,
    *,
    load_type_options: bool = True,
) -> dict[str, object]:
    """Build the shared filtered and ordered query used by the UI and exports."""
    query_text = (values.get("q") or "").strip()
    selected_status = (values.get("status") or "all").strip()
    if selected_status not in {"all", "included", "excluded"}:
        selected_status = "all"
    selected_type = (values.get("type") or "").strip()
    selected_validation = (values.get("validation") or "all").strip()
    if selected_validation not in {"all", "validated", "unvalidated"}:
        selected_validation = "all"
    selected_source = (values.get("source") or "all").strip()
    if selected_source not in {"all", "xlsx"}:
        selected_source = "all"
    selected_sort = (values.get("sort") or "year").strip().lower()
    if selected_sort not in {
        "title", "creator", "year", "type", "affiliation", "validation", "status"
    }:
        selected_sort = "year"
    selected_direction = (values.get("direction") or "desc").strip().lower()
    if selected_direction not in {"asc", "desc"}:
        selected_direction = "desc"

    works_query, included_expr, _datestamp_expr = institutional_work_query(config)
    if query_text:
        pattern = f"%{query_text}%"
        works_query = works_query.filter(or_(
            CanonicalWork.title.ilike(pattern),
            CanonicalWork.doi_normalized.ilike(pattern),
            WorkCache.title.ilike(pattern),
            WorkCache.journal_title.ilike(pattern),
            OpenAlexWorkMetadata.author_names.ilike(pattern),
            OpenAlexWorkMetadata.institution_names.ilike(pattern),
            OpenAlexWorkMetadata.institution_rors.ilike(pattern),
        ))
    if selected_status == "included":
        works_query = works_query.filter(included_expr.is_(True))
    elif selected_status == "excluded":
        works_query = works_query.filter(included_expr.is_(False))
    if selected_validation == "validated":
        works_query = works_query.filter(
            OpenAlexInstitutionWorkFact.has_selected_affiliation.is_(True)
        )
    elif selected_validation == "unvalidated":
        works_query = works_query.filter(or_(
            OpenAlexInstitutionWorkFact.id.is_(None),
            OpenAlexInstitutionWorkFact.has_selected_affiliation.is_(False),
        ))
    if selected_source == "xlsx":
        works_query = works_query.filter(
            OaiPmhWorkSelection.is_included.is_(True),
            OaiPmhWorkSelection.decision_source == "xlsx",
        )

    type_expr = func.coalesce(OpenAlexWorkMetadata.type, WorkCache.type)
    type_options = []
    if load_type_options:
        type_options = [
            value
            for (value,) in works_query.with_entities(type_expr)
            .filter(type_expr.isnot(None), type_expr != "")
            .distinct()
            .order_by(type_expr.asc())
            .all()
        ]
    if selected_type and (not load_type_options or selected_type in type_options):
        works_query = works_query.filter(type_expr == selected_type)
    elif selected_type:
        selected_type = ""

    sort_expressions = {
        "title": func.lower(func.coalesce(
            OpenAlexWorkMetadata.title,
            CanonicalWork.title,
            WorkCache.title,
        )),
        "creator": func.lower(func.coalesce(
            OpenAlexWorkMetadata.author_names,
            ResearcherCache.credit_name,
            ResearcherCache.family_name,
            WorkCache.orcid,
        )),
        "year": func.coalesce(
            OpenAlexWorkMetadata.publication_year,
            CanonicalWork.publication_year,
        ),
        "type": func.lower(type_expr),
        "affiliation": func.lower(OpenAlexWorkMetadata.institution_names),
        "validation": OpenAlexInstitutionWorkFact.has_selected_affiliation,
        "status": included_expr,
    }
    sort_expression = sort_expressions[selected_sort]
    ordering = (
        sort_expression.asc().nullslast()
        if selected_direction == "asc"
        else sort_expression.desc().nullslast()
    )
    return {
        "query": works_query,
        "ordering": ordering,
        "query_text": query_text,
        "status": selected_status,
        "type": selected_type,
        "validation": selected_validation,
        "source": selected_source,
        "sort": selected_sort,
        "direction": selected_direction,
        "type_options": type_options,
    }


def _article_export_headers() -> list[str]:
    """Return localized, stable columns for an institutional audit export."""
    return [
        _("Article"),
        _("DOI"),
        _("Creators"),
        _("Year"),
        _("Type"),
        _("Journal"),
        _("Affiliations"),
        _("Affiliation RORs"),
        _("Selected institution affiliation"),
        _("OpenAlex validation"),
        _("OAI status"),
        _("Activation origin"),
        _("ORCID"),
        _("Article URL"),
    ]


def _iter_article_export_rows(query, config, ror_id: str):
    """Yield full audit rows while enriching affiliations in bounded chunks."""
    pending = []
    for query_row in query.yield_per(250):
        pending.append(work_from_query_row(
            query_row,
            config,
            include_metadata=False,
        ))
        if len(pending) >= 250:
            attach_institutional_affiliations(pending, ror_id)
            for record in pending:
                yield _article_export_row(record)
            pending = []
    if pending:
        attach_institutional_affiliations(pending, ror_id)
        for record in pending:
            yield _article_export_row(record)


def _article_export_row(record) -> list[object]:
    """Serialize one institutional work with every affiliation for auditing."""
    if record.publication_source == "xlsx":
        publication_source = _("Excel import")
    elif record.has_manual_override:
        publication_source = _("Manual decision")
    else:
        publication_source = _("Automatic policy")
    return [
        record.title or _("Untitled article"),
        record.doi_normalized or "",
        "; ".join(record.creators),
        record.publication_year or "",
        record.document_type or "",
        record.journal_title or "",
        "; ".join(item["name"] for item in record.affiliations),
        "; ".join(
            item["ror_id"] for item in record.affiliations if item.get("ror_id")
        ),
        "; ".join(
            item["name"] for item in record.affiliations if item.get("is_selected")
        ),
        _("Validated") if record.is_openalex_validated else _("Not validated"),
        _("Exposed") if record.is_included else _("Excluded"),
        publication_source,
        record.orcid or "",
        record.landing_page_url or "",
    ]


def _safe_spreadsheet_cell(value, *, xlsx: bool = False):
    """Remove unsafe controls and neutralize spreadsheet formulas in text."""
    if not isinstance(value, str):
        return value
    safe_value = EXCEL_ILLEGAL_CHARACTERS_RE.sub("", value)
    if xlsx:
        safe_value = safe_value[:32767]
    if safe_value.lstrip().startswith(("=", "+", "-", "@")):
        safe_value = "'" + safe_value
    return safe_value


def _read_doi_import(upload) -> dict[str, object]:
    """Validate one XLSX upload and return normalized DOI row statistics."""
    stream = upload.stream
    try:
        stream.seek(0, 2)
        upload_size = stream.tell()
        stream.seek(0)
    except (AttributeError, OSError):
        upload_size = request.content_length or 0
    if upload_size > DOI_IMPORT_MAX_BYTES:
        raise DoiImportError(_("The DOI spreadsheet cannot exceed 5 MB."))

    workbook = load_workbook(stream, read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            raise DoiImportError(_("The DOI spreadsheet does not contain any worksheets."))
        worksheet = workbook["DOI"] if "DOI" in workbook.sheetnames else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header = None
        header_row_number = 0
        for header_row_number, candidate in enumerate(rows, start=1):
            if any(str(value or "").strip() for value in candidate):
                header = candidate
                break
            if header_row_number >= 20:
                break
        if header is None:
            raise DoiImportError(_("The DOI spreadsheet is empty."))

        normalized_headers = [
            str(value or "").strip().lower().replace("\ufeff", "")
            for value in header
        ]
        if "doi" not in normalized_headers:
            raise DoiImportError(_("The spreadsheet must contain a column named doi."))
        doi_column = normalized_headers.index("doi")

        submitted_count = 0
        invalid_count = 0
        duplicate_count = 0
        normalized_rows = []
        seen = set()
        for data_offset, row in enumerate(rows, start=1):
            if data_offset > DOI_IMPORT_MAX_ROWS:
                raise DoiImportError(_(
                    "The DOI spreadsheet can contain at most %(count)s data rows.",
                    count=DOI_IMPORT_MAX_ROWS,
                ))
            value = row[doi_column] if doi_column < len(row) else None
            raw_doi = str(value or "").strip()
            if not raw_doi:
                continue
            submitted_count += 1
            normalized = normalize_doi(raw_doi)
            if not normalized:
                invalid_count += 1
                continue
            normalized_rows.append(normalized)
            if normalized in seen:
                duplicate_count += 1
            else:
                seen.add(normalized)

        if submitted_count == 0:
            raise DoiImportError(_("Enter at least one DOI in the spreadsheet."))
        return {
            "submitted_count": submitted_count,
            "invalid_count": invalid_count,
            "duplicate_count": duplicate_count,
            "normalized_rows": normalized_rows,
            "header_row_number": header_row_number,
        }
    finally:
        workbook.close()


@bp_oai_pmh.route("/oai-pmh/stylesheet/<language>.xsl")
def browser_stylesheet(language: str):
    """Translate browser presentation without changing harvested metadata."""
    if language not in DEFAULT_LANGUAGES:
        abort(404)
    with force_locale(language):
        content = render_template("oai_pmh/provider.xsl", language=language)
    return Response(content, content_type="text/xsl; charset=utf-8")


@bp_oai_pmh.route("/oai/<public_key>", methods=["GET", "POST"], defaults={"harvester_key": None})
@bp_oai_pmh.route("/oai/<public_key>/<harvester_key>", methods=["GET", "POST"])
@csrf.exempt
def provider(public_key: str, harvester_key: str | None):
    """Authorize the requested institution and credential before building metadata."""
    headers = {"Cache-Control": "private, no-store", "Referrer-Policy": "no-referrer"}
    config = OaiPmhInstitutionConfig.query.filter_by(
        public_key=public_key,
        provider_enabled=True,
    ).first()
    if not config:
        return Response(_("OAI-PMH repository not found."), status=404, mimetype="text/plain", headers=headers)
    if harvester_key is not None:
        harvester = OaiPmhHarvester.query.filter_by(
            config_id=config.id, access_key=harvester_key, is_enabled=True,
        ).first()
        if not harvester:
            return Response(_("OAI-PMH repository not found."), status=404, mimetype="text/plain", headers=headers)
    elif config.harvester_access_restricted:
        return Response(_("A private harvesting URL is required."), status=403, mimetype="text/plain", headers=headers)

    params = CombinedMultiDict((request.args, request.form))
    language = str(get_locale())
    if language not in DEFAULT_LANGUAGES:
        language = "en"
    payload = build_oai_response(
        config, params, _provider_url(config.public_key, harvester_key),
        stylesheet_url=url_for("oai_pmh.browser_stylesheet", language=language),
    )
    return Response(payload, status=200, content_type="text/xml; charset=utf-8", headers=headers)


def _provider_url(public_key: str, harvester_key: str | None = None) -> str:
    configured_base = (current_app.config.get("APP_BASE_URL") or "").rstrip("/")
    values = {"public_key": public_key, "harvester_key": harvester_key}
    path = url_for("oai_pmh.provider", **values)
    return f"{configured_base}{path}" if configured_base else url_for(
        "oai_pmh.provider", **values, _external=True
    )


def _user_display_names(user_ids) -> dict[int, str]:
    """Resolve audit actor labels without exposing account identifiers."""
    if not user_ids:
        return {}
    users = User.query.filter(User.id.in_(user_ids)).all()
    return {
        user.id: user.full_name or user.username or _("Unknown user")
        for user in users
    }


def _return_to_articles():
    return url_for(
        "oai_pmh.articles",
        q=(request.form.get("q") or "").strip(),
        status=(request.form.get("status") or "all").strip(),
        type=(request.form.get("type") or "").strip(),
        validation=(request.form.get("validation") or "all").strip(),
        source=(request.form.get("source") or "all").strip(),
        sort=(request.form.get("sort") or "year").strip(),
        direction=(request.form.get("direction") or "desc").strip(),
        page=max(_safe_int(request.form.get("page"), 1), 1),
    )


def _return_to_doi_import():
    return url_for("oai_pmh.doi_import_workspace")


def _can_manage_provider() -> bool:
    return bool(session.get("is_admin") or session.get("is_manager"))


def _can_manage_oai_content() -> bool:
    return bool(
        session.get("is_admin")
        or session.get("is_manager")
        or session.get("is_oai_user")
    )


def _metadata_catalog_view(mapping: dict[str, str]) -> list[dict]:
    """Return translated field definitions for the mapping editor."""
    labels = {
        "title": _("Title"),
        "creator": _("Creator"),
        "subject": _("Subject"),
        "description": _("Description"),
        "publisher": _("Publisher"),
        "contributor": _("Contributor"),
        "date": _("Date"),
        "type": _("Type"),
        "format": _("Format"),
        "identifier": _("Identifier"),
        "source": _("Source"),
        "language": _("Language"),
        "relation": _("Relation"),
        "coverage": _("Coverage"),
        "rights": _("Rights"),
        "publication_year": _("Publication year"),
        "journal_title": _("Journal title"),
        "volume": _("Volume"),
        "issue": _("Issue"),
        "first_page": _("First page"),
        "last_page": _("Last page"),
        "corresponding_creator": _("Corresponding creator"),
        "doi": _("DOI"),
        "openalex_id": _("OpenAlex identifier"),
        "creator_orcid": _("Creator ORCID"),
        "issn": _("ISSN"),
        "pmid": _("PMID"),
        "pmcid": _("PMCID"),
        "source_id": _("Source OpenAlex identifier"),
        "landing_page_url": _("Landing page URL"),
        "pdf_url": _("PDF URL"),
        "license": _("License"),
        "access_right": _("Access right"),
        "oa_status": _("Open-access status"),
        "is_open_access": _("Is open access"),
        "source_type": _("Source type"),
        "indexed_in": _("Indexed in"),
        "institution": _("Institution"),
        "institution_ror": _("Institution ROR"),
        "country": _("Country"),
        "funder": _("Funder"),
        "award": _("Award or grant"),
        "sdg": _("Sustainable Development Goal"),
        "keyword": _("Keyword"),
        "topic": _("Topic"),
        "primary_topic": _("Primary topic"),
        "topic_field": _("Topic field"),
        "topic_domain": _("Topic domain"),
        "cited_by_count": _("Citation count"),
        "fwci": _("Field-weighted citation impact (FWCI)"),
    }
    return [
        {
            "key": field,
            "label": labels[field],
            "default_target": target,
            "group": group,
            "active": field in mapping,
            "target": mapping.get(field, target),
            "required": required,
        }
        for field, target, group, _enabled_by_default, required in METADATA_FIELD_CATALOG
    ]


def _metadata_group_labels() -> dict[str, str]:
    return {
        "core": _("Dublin Core base"),
        "publication": _("Publication details"),
        "identifier": _("Identifiers and links"),
        "access": _("Access and source"),
        "context": _("Affiliations and funding"),
        "enrichment": _("OpenAlex enrichment"),
    }


def _safe_int(value, default: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default
