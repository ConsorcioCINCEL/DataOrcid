"""Regression tests for durable, private large-file exports."""

import hashlib
import json
import os
import tempfile
import unittest
from datetime import timedelta
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from app import create_app, db
from app.models import SyncJob, SyncJobStep, User, WorkCache, utc_now
from app.services.background_jobs import run_queued_job
from app.services.export_jobs import cleanup_expired_exports


class BackgroundExportTest(unittest.TestCase):
    def setUp(self):
        self.export_root = tempfile.TemporaryDirectory()
        self.app = create_app({
            "TESTING": True,
            "SECRET_KEY": "test-secret",
            "SQLALCHEMY_DATABASE_URI": "sqlite://",
            "JOB_EXECUTION_MODE": "queue",
            "EXPORT_DIRECTORY": self.export_root.name,
            "EXPORT_RETENTION_HOURS": 24,
            "WTF_CSRF_ENABLED": False,
        })
        with self.app.app_context():
            db.create_all()
            owner = User(username="owner@example.test", ror_id="01test123")
            owner.set_password("test-password")
            outsider = User(username="outsider@example.test", ror_id="02other456")
            outsider.set_password("test-password")
            db.session.add_all([owner, outsider])
            db.session.flush()
            db.session.add(WorkCache(
                ror_id="01test123",
                orcid="0000-0001-2345-6789",
                title="A stable export row",
                type="journal-article",
                put_code=7,
                pub_year="2026",
                doi="10.1234/example",
                visibility="public",
            ))
            db.session.commit()
            self.owner_id = owner.id
            self.outsider_id = outsider.id

        self.client = self.app.test_client()
        self._login(self.client, self.owner_id, "01test123")

    def tearDown(self):
        with self.app.app_context():
            engine = db.engine
            db.session.remove()
            db.drop_all()
            engine.dispose()
        self.export_root.cleanup()

    @staticmethod
    def _login(client, user_id, ror_id):
        with client.session_transaction() as client_session:
            client_session.update(
                logged_in=True,
                user_id=user_id,
                username="user@example.test",
                ror_id=ror_id,
                is_admin=False,
                is_manager=False,
            )

    def test_background_csv_is_private_and_matches_direct_export(self):
        direct = self.client.get("/download/all-works/cache")
        queued_response = self.client.get(
            "/download/all-works/cache?background=1",
            headers={"Accept": "application/json"},
        )

        self.assertEqual(200, direct.status_code)
        self.assertEqual(202, queued_response.status_code)
        job_id = queued_response.get_json()["job"]["id"]
        self.assertEqual(job_id, run_queued_job(self.app, "export-test-worker"))

        status = self.client.get(f"/exports/jobs/{job_id}")
        payload = status.get_json()["job"]
        self.assertEqual("success", payload["status"])
        self.assertEqual(1, payload["records"])
        self.assertTrue(payload["download_url"].endswith(f"/{job_id}/download"))

        generated = self.client.get(payload["download_url"])
        self.assertEqual(200, generated.status_code)
        self.assertEqual(direct.data, generated.data)
        self.assertIn("private, no-store", generated.headers["Cache-Control"])
        self.assertEqual(1, len(list(Path(self.export_root.name).glob("*.csv"))))

        outsider = self.app.test_client()
        self._login(outsider, self.outsider_id, "02other456")
        self.assertEqual(404, outsider.get(f"/exports/jobs/{job_id}").status_code)
        self.assertEqual(404, outsider.get(payload["download_url"]).status_code)

    def test_duplicate_click_reuses_one_queued_export(self):
        first_payload = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()
        second_payload = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()
        first = first_payload["job"]
        second = second_payload["job"]

        self.assertEqual(first["id"], second["id"])
        self.assertFalse(first_payload["reused"])
        self.assertTrue(second_payload["reused"])
        with self.app.app_context():
            self.assertEqual(1, SyncJob.query.filter_by(job_type="export").count())

    def test_completed_identical_export_reuses_the_existing_private_file(self):
        first_payload = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()
        first = first_payload["job"]
        self.assertEqual(first["id"], run_queued_job(self.app, "reuse-test-worker"))
        first_status = self.client.get(
            f"/exports/jobs/{first['id']}"
        ).get_json()["job"]
        original_file = Path(self.export_root.name) / f"{first['id']}.csv"
        original_bytes = original_file.read_bytes()

        second_payload = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()

        self.assertTrue(second_payload["reused"])
        self.assertEqual(first["id"], second_payload["job"]["id"])
        self.assertEqual(first_status["download_url"], second_payload["job"]["download_url"])
        self.assertEqual(original_bytes, original_file.read_bytes())
        with self.app.app_context():
            self.assertEqual(1, SyncJob.query.filter_by(job_type="export").count())
        self.assertEqual(1, len(list(Path(self.export_root.name).glob("*.csv"))))

    def test_valid_legacy_job_is_adopted_without_generating_a_second_file(self):
        first = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()["job"]
        self.assertEqual(first["id"], run_queued_job(self.app, "legacy-file-worker"))

        with self.app.app_context():
            job = db.session.get(SyncJob, first["id"])
            payload = json.loads(json.dumps(job.payload_json))
            parameters = payload["args"][2]
            parameters.pop("_source_revision")
            legacy_signature = hashlib.sha256(
                json.dumps(parameters, sort_keys=True).encode("utf-8")
            ).hexdigest()[:16]
            job.name = (
                f"export-{self.owner_id}-institution_works-csv-{legacy_signature}"
            )
            job.payload_json = payload
            db.session.commit()

        second_payload = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()

        self.assertTrue(second_payload["reused"])
        self.assertEqual(first["id"], second_payload["job"]["id"])
        with self.app.app_context():
            adopted = db.session.get(SyncJob, first["id"])
            self.assertIn("_source_revision", adopted.payload_json["args"][2])
            self.assertEqual(1, SyncJob.query.filter_by(job_type="export").count())
        self.assertEqual(1, len(list(Path(self.export_root.name).glob("*.csv"))))

    def test_completed_export_is_never_reused_by_another_account(self):
        owner_job = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()["job"]
        self.assertEqual(owner_job["id"], run_queued_job(self.app, "owner-file-worker"))

        with self.app.app_context():
            db.session.get(User, self.outsider_id).ror_id = "01test123"
            db.session.commit()
        outsider = self.app.test_client()
        self._login(outsider, self.outsider_id, "01test123")
        outsider_payload = outsider.get(
            "/download/all-works/cache?background=1"
        ).get_json()

        self.assertFalse(outsider_payload["reused"])
        self.assertNotEqual(owner_job["id"], outsider_payload["job"]["id"])
        with self.app.app_context():
            self.assertEqual(2, SyncJob.query.filter_by(job_type="export").count())

    def test_source_change_invalidates_a_completed_export(self):
        first = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()["job"]
        self.assertEqual(first["id"], run_queued_job(self.app, "revision-test-worker"))

        with self.app.app_context():
            changed_at = utc_now()
            db.session.add(SyncJob(
                id="00000000-0000-0000-0000-000000000099",
                name="source-data-refresh",
                job_type="full_institution_sync",
                requested_by_user_id=self.owner_id,
                status="success",
                created_at=changed_at,
                finished_at=changed_at,
            ))
            db.session.commit()

        second_payload = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()

        self.assertFalse(second_payload["reused"])
        self.assertNotEqual(first["id"], second_payload["job"]["id"])
        with self.app.app_context():
            self.assertEqual(2, SyncJob.query.filter_by(job_type="export").count())

    def test_missing_completed_file_is_replaced_instead_of_reused(self):
        first = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()["job"]
        self.assertEqual(first["id"], run_queued_job(self.app, "missing-file-worker"))
        (Path(self.export_root.name) / f"{first['id']}.csv").unlink()

        second_payload = self.client.get(
            "/download/all-works/cache?background=1"
        ).get_json()

        self.assertFalse(second_payload["reused"])
        self.assertNotEqual(first["id"], second_payload["job"]["id"])
        with self.app.app_context():
            self.assertIsNone(db.session.get(SyncJob, first["id"]))
            self.assertEqual(1, SyncJob.query.filter_by(job_type="export").count())

    def test_background_xlsx_is_a_valid_workbook(self):
        with self.app.app_context():
            work = WorkCache.query.one()
            work.title = "A stable\x02 export row"
            db.session.commit()

        queued = self.client.get(
            "/download/all-works/cache?format=excel&background=1"
        ).get_json()["job"]
        self.assertEqual(queued["id"], run_queued_job(self.app, "xlsx-export-worker"))
        status = self.client.get(f"/exports/jobs/{queued['id']}").get_json()["job"]
        generated = self.client.get(status["download_url"])

        self.assertEqual("success", status["status"])
        workbook = load_workbook(BytesIO(generated.data), read_only=True)
        try:
            self.assertEqual(["Works"], workbook.sheetnames)
            rows = list(workbook["Works"].iter_rows(values_only=True))
            self.assertEqual("A stable export row", rows[1][3])
        finally:
            workbook.close()

    def test_recent_export_list_is_scoped_to_the_requester(self):
        job = self.client.get("/download/all-works/cache?background=1").get_json()["job"]
        listing = self.client.get("/exports/jobs").get_json()["jobs"]
        self.assertEqual([job["id"]], [item["id"] for item in listing])

        outsider = self.app.test_client()
        self._login(outsider, self.outsider_id, "02other456")
        self.assertEqual([], outsider.get("/exports/jobs").get_json()["jobs"])

    def test_filtered_researcher_directory_export_keeps_the_same_csv(self):
        url = "/researcher-list/export?format=csv&q=0000-0001"
        direct = self.client.get(url)
        queued = self.client.get(f"{url}&background=1").get_json()["job"]

        self.assertEqual(200, direct.status_code)
        self.assertEqual(queued["id"], run_queued_job(self.app, "directory-export-worker"))
        status = self.client.get(f"/exports/jobs/{queued['id']}").get_json()["job"]
        self.assertEqual("success", status["status"])
        generated = self.client.get(status["download_url"])
        self.assertEqual(direct.data, generated.data)

    def test_duplicate_report_background_csv_matches_direct_export(self):
        url = "/duplicates/download?scope=current&format=csv"
        direct = self.client.get(url)
        queued = self.client.get(f"{url}&background=1").get_json()["job"]

        self.assertEqual(queued["id"], run_queued_job(self.app, "duplicate-export-worker"))
        status = self.client.get(f"/exports/jobs/{queued['id']}").get_json()["job"]
        self.assertEqual("success", status["status"])
        generated = self.client.get(status["download_url"])
        self.assertEqual(direct.data, generated.data)

    def test_expired_cleanup_removes_files_jobs_and_steps_but_keeps_active_work(self):
        old = utc_now() - timedelta(hours=30)
        recent = utc_now() - timedelta(hours=1)
        expired_id = "00000000-0000-0000-0000-000000000001"
        failed_id = "00000000-0000-0000-0000-000000000002"
        active_id = "00000000-0000-0000-0000-000000000003"
        recent_id = "00000000-0000-0000-0000-000000000004"
        export_root = Path(self.export_root.name)

        expired_file = export_root / f"{expired_id}.csv"
        active_file = export_root / f"{active_id}.csv"
        recent_file = export_root / f"{recent_id}.csv"
        orphan_file = export_root / "orphan.csv"
        for path in (expired_file, active_file, recent_file, orphan_file):
            path.write_text("content", encoding="utf-8")
            os.utime(path, (old.timestamp(), old.timestamp()))

        with self.app.app_context():
            db.session.add_all([
                SyncJob(
                    id=expired_id,
                    name="expired-success",
                    job_type="export",
                    requested_by_user_id=self.owner_id,
                    status="success",
                    created_at=old,
                    finished_at=old,
                    result_json={"export": {"stored_name": expired_file.name}},
                ),
                SyncJob(
                    id=failed_id,
                    name="expired-failure",
                    job_type="export",
                    requested_by_user_id=self.owner_id,
                    status="failed",
                    created_at=old,
                    finished_at=old,
                ),
                SyncJob(
                    id=active_id,
                    name="old-active",
                    job_type="export",
                    requested_by_user_id=self.owner_id,
                    status="running",
                    created_at=old,
                    result_json={"export": {"stored_name": active_file.name}},
                ),
                SyncJob(
                    id=recent_id,
                    name="recent-success",
                    job_type="export",
                    requested_by_user_id=self.owner_id,
                    status="success",
                    created_at=recent,
                    finished_at=recent,
                    result_json={"export": {"stored_name": recent_file.name}},
                ),
            ])
            db.session.add(SyncJobStep(
                sync_job_id=expired_id,
                name="prepare_file",
                status="success",
            ))
            db.session.commit()

            result = cleanup_expired_exports(force=True)

            self.assertEqual({"files": 2, "jobs": 2}, result)
            self.assertIsNone(db.session.get(SyncJob, expired_id))
            self.assertIsNone(db.session.get(SyncJob, failed_id))
            self.assertIsNone(
                SyncJobStep.query.filter_by(sync_job_id=expired_id).first()
            )
            self.assertIsNotNone(db.session.get(SyncJob, active_id))
            self.assertIsNotNone(db.session.get(SyncJob, recent_id))
        self.assertFalse(expired_file.exists())
        self.assertFalse(orphan_file.exists())
        self.assertTrue(active_file.exists())
        self.assertTrue(recent_file.exists())

    def test_clear_all_deletes_only_current_users_finished_exports(self):
        owner_success_id = "00000000-0000-0000-0000-000000000011"
        owner_failed_id = "00000000-0000-0000-0000-000000000012"
        owner_active_id = "00000000-0000-0000-0000-000000000013"
        outsider_id = "00000000-0000-0000-0000-000000000014"
        export_root = Path(self.export_root.name)
        owner_file = export_root / f"{owner_success_id}.xlsx"
        owner_temporary = export_root / f".{owner_failed_id}.csv.tmp"
        outsider_file = export_root / f"{outsider_id}.csv"
        for path in (owner_file, owner_temporary, outsider_file):
            path.write_text("content", encoding="utf-8")

        with self.app.app_context():
            db.session.add_all([
                SyncJob(
                    id=owner_success_id,
                    name="owner-success",
                    job_type="export",
                    requested_by_user_id=self.owner_id,
                    status="success",
                    finished_at=utc_now(),
                    result_json={"export": {"stored_name": owner_file.name}},
                ),
                SyncJob(
                    id=owner_failed_id,
                    name="owner-failed",
                    job_type="export",
                    requested_by_user_id=self.owner_id,
                    status="failed",
                    finished_at=utc_now(),
                ),
                SyncJob(
                    id=owner_active_id,
                    name="owner-active",
                    job_type="export",
                    requested_by_user_id=self.owner_id,
                    status="running",
                ),
                SyncJob(
                    id=outsider_id,
                    name="outsider-success",
                    job_type="export",
                    requested_by_user_id=self.outsider_id,
                    status="success",
                    finished_at=utc_now(),
                    result_json={"export": {"stored_name": outsider_file.name}},
                ),
            ])
            db.session.add(SyncJobStep(
                sync_job_id=owner_success_id,
                name="prepare_file",
                status="success",
            ))
            db.session.commit()

        response = self.client.post("/exports/jobs/clear")

        self.assertEqual(200, response.status_code)
        self.assertEqual(
            {"deleted_files": 2, "deleted_jobs": 2},
            response.get_json(),
        )
        with self.app.app_context():
            self.assertIsNone(db.session.get(SyncJob, owner_success_id))
            self.assertIsNone(db.session.get(SyncJob, owner_failed_id))
            self.assertIsNotNone(db.session.get(SyncJob, owner_active_id))
            self.assertIsNotNone(db.session.get(SyncJob, outsider_id))
            self.assertIsNone(
                SyncJobStep.query.filter_by(sync_job_id=owner_success_id).first()
            )
        self.assertFalse(owner_file.exists())
        self.assertFalse(owner_temporary.exists())
        self.assertTrue(outsider_file.exists())


if __name__ == "__main__":
    unittest.main()
