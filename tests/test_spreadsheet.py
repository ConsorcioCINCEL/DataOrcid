"""Regression tests for values written to XLSX exports."""

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import load_workbook

from app.services.export_jobs import _write_dict_rows
from app.spreadsheet import EXCEL_CELL_TEXT_LIMIT, excel_cell, excel_safe_dataframe


class SpreadsheetSafetyTest(unittest.TestCase):
    def test_excel_cell_removes_xml_controls_but_preserves_layout_whitespace(self):
        value = "before\x00\x02\x0b\x0c\x1f\tafter\nnext\rline\ud800\ufffe"

        self.assertEqual("before\tafter\nnext\rline", excel_cell(value))

    def test_excel_cell_truncates_to_the_openpyxl_limit(self):
        value = excel_cell("x" * (EXCEL_CELL_TEXT_LIMIT + 100))

        self.assertEqual(EXCEL_CELL_TEXT_LIMIT, len(value))
        self.assertTrue(value.endswith("... [truncated for Excel]"))

    def test_excel_safe_dataframe_only_changes_text_values(self):
        source = pd.DataFrame({
            "title": ["invalid\x02title"],
            "count": [7],
            "available": [True],
        })

        safe = excel_safe_dataframe(source)

        self.assertEqual("invalidtitle", safe.loc[0, "title"])
        self.assertEqual(7, safe.loc[0, "count"])
        self.assertTrue(safe.loc[0, "available"])
        self.assertEqual("invalid\x02title", source.loc[0, "title"])

    def test_streaming_xlsx_writer_uses_the_shared_cell_sanitizer(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "streamed.xlsx"
            count = _write_dict_rows(
                iter([{"title": "streamed\x02title", "count": 3}]),
                ["title", "count"],
                path,
                "excel",
                "Rows",
                None,
            )

            workbook = load_workbook(path, read_only=True)
            try:
                rows = list(workbook["Rows"].iter_rows(values_only=True))
            finally:
                workbook.close()

        self.assertEqual(1, count)
        self.assertEqual(("streamedtitle", 3), rows[1])


if __name__ == "__main__":
    unittest.main()
