"""OAI-PMH provider, article-selection, and tenant-isolation tests."""

import csv
from datetime import timedelta
from io import BytesIO, StringIO
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET

from flask import Flask
from lxml import etree
from openpyxl import Workbook, load_workbook

from app import babel, csrf, db
from app.blueprints.admin import bp_admin
from app.blueprints.oai_pmh import bp_oai_pmh
from app.models import (
    CanonicalWork,
    InstitutionRegistry,
    OaiPmhDoiImportBatch,
    OaiPmhDoiImportChange,
    OaiPmhInstitutionConfig,
    OaiPmhWorkSelection,
    OpenAlexInstitutionWorkFact,
    OpenAlexWorkInstitution,
    OpenAlexWorkMetadata,
    User,
    WorkCache,
    WorkRecordLink,
    utc_now,
)
from app.services.oai_pmh_service import OAI_NS, oai_repository_summaries


class OaiPmhModuleTest(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(
            SECRET_KEY="test-key",
            SQLALCHEMY_DATABASE_URI="sqlite://",
            SQLALCHEMY_TRACK_MODIFICATIONS=False,
            OAI_PROVIDER_PAGE_SIZE=1,
            OAI_RESUMPTION_TOKEN_MAX_AGE=3600,
            WTF_CSRF_ENABLED=False,
            TESTING=True,
        )
        db.init_app(self.app)
        babel.init_app(self.app)
        csrf.init_app(self.app)
        self.app.register_blueprint(bp_oai_pmh)
        self.app.register_blueprint(bp_admin)
        self.app.add_url_rule("/login", endpoint="auth.login", view_func=lambda: "login")
        self.app.add_url_rule("/", endpoint="main.index", view_func=lambda: "home")

        with self.app.app_context():
            db.create_all()
            manager = User(
                username="manager@example.org",
                email="manager@example.org",
                is_manager=True,
                ror_id="01aaa1111",
            )
            manager.set_password("test-password")
            standard = User(username="user@example.org", ror_id="01aaa1111")
            standard.set_password("test-password")
            oai_editor = User(
                username="oai@example.org",
                ror_id="01aaa1111",
                is_oai_user=True,
            )
            oai_editor.set_password("test-password")
            admin = User(
                username="admin@example.org",
                email="admin@example.org",
                is_admin=True,
            )
            admin.set_password("test-password")
            db.session.add_all([
                manager,
                standard,
                oai_editor,
                admin,
                InstitutionRegistry(ror_id="01aaa1111", name="University A"),
                InstitutionRegistry(ror_id="02bbb2222", name="University B"),
            ])
            now = utc_now().replace(microsecond=0) - timedelta(minutes=5)
            config_a = OaiPmhInstitutionConfig(
                ror_id="01aaa1111",
                public_key="a" * 48,
                provider_enabled=True,
                repository_name="DataORCID — University A",
                admin_email="oai-a@example.org",
                publication_policy="validated",
                policy_updated_at=now,
                created_at=now,
                updated_at=now,
            )
            config_b = OaiPmhInstitutionConfig(
                ror_id="02bbb2222",
                public_key="b" * 48,
                provider_enabled=True,
                repository_name="DataORCID — University B",
                admin_email="oai-b@example.org",
                publication_policy="validated",
                policy_updated_at=now,
                created_at=now,
                updated_at=now,
            )
            db.session.add_all([config_a, config_b])
            db.session.flush()

            work_a = self._add_work(
                "01aaa1111", "0000-0001-0000-0001", 101, "Article A",
                "10.1234/a", "a", now,
                validated=True,
            )
            work_a2 = self._add_work(
                "01aaa1111", "0000-0001-0000-0002", 102, "Article A2",
                "10.1234/a2", "b", now + timedelta(seconds=1),
                validated=False,
            )
            work_b = self._add_work(
                "02bbb2222", "0000-0002-0000-0001", 201, "Article B",
                "10.1234/b", "c", now + timedelta(seconds=2),
                validated=True,
            )
            db.session.commit()
            self.manager_id = manager.id
            self.standard_id = standard.id
            self.oai_editor_id = oai_editor.id
            self.admin_id = admin.id
            self.work_a_id = work_a.id
            self.work_a2_id = work_a2.id
            self.work_b_id = work_b.id
            self.key_a = work_a.canonical_key
            self.key_a2 = work_a2.canonical_key

        self.client = self.app.test_client()

    def tearDown(self):
        with self.app.app_context():
            engine = db.engine
            db.session.remove()
            db.drop_all()
            db.session.remove()
            engine.dispose()

    @staticmethod
    def _add_work(ror_id, orcid, put_code, title, doi, suffix, updated_at, *, validated):
        local = WorkCache(
            ror_id=ror_id,
            orcid=orcid,
            put_code=put_code,
            title=title,
            doi=doi,
            pub_year="2025",
            type="journal-article",
            journal_title="Test Journal",
            url=f"https://example.org/{suffix}",
            visibility="public",
            created_at=updated_at,
        )
        db.session.add(local)
        db.session.flush()
        canonical = CanonicalWork(
            canonical_key=f"doi:{suffix * 64}",
            doi_normalized=doi,
            title=title,
            title_normalized=title.lower(),
            publication_year=2025,
            record_count=1,
            created_at=updated_at,
            updated_at=updated_at,
        )
        db.session.add(canonical)
        db.session.flush()
        db.session.add(WorkRecordLink(
            canonical_work_id=canonical.id,
            work_cache_id=local.id,
            ror_id=ror_id,
            orcid=orcid,
            source_record_key=f"put:{put_code}",
            created_at=updated_at,
        ))
        db.session.add(OpenAlexInstitutionWorkFact(
            ror_id=ror_id,
            openalex_cache_key=doi,
            representative_work_cache_id=local.id,
            source_record_count=1,
            has_valid_doi=True,
            has_local_title=True,
            has_selected_affiliation=validated,
            refreshed_at=updated_at,
        ))
        db.session.add(OpenAlexWorkMetadata(
            doi_normalized=doi,
            openalex_id=f"W{put_code}",
            title=title,
            publication_year=2025,
            publication_date="2025-04-01",
            type="article",
            language="en",
            author_names="Author One; Author Two",
            keywords="topic-id | Research integrity | 0.9000 || metadata | Metadata | 0.8000",
            source_name="Enriched Journal",
            source_issn_l="1234-5678",
            volume="42",
            issue="7",
            first_page="10",
            last_page="22",
            institution_names=(
                "University A; Partner University"
                if ror_id == "01aaa1111"
                else "University B; Partner University"
            ),
            institution_rors=ror_id,
            countries="CL",
            primary_landing_page_url=f"https://doi.org/{doi}",
            created_at=updated_at,
            updated_at=updated_at,
            fetched_at=updated_at,
        ))
        db.session.add_all([
            OpenAlexWorkInstitution(
                doi_normalized=doi,
                openalex_id=f"W{put_code}",
                institution_id=f"I{put_code}",
                institution_name=(
                    "University A" if ror_id == "01aaa1111" else "University B"
                ),
                ror_id=ror_id,
                country_code="CL",
                author_count=1,
                has_corresponding_author=True,
                created_at=updated_at,
            ),
            OpenAlexWorkInstitution(
                doi_normalized=doi,
                openalex_id=f"W{put_code}",
                institution_id=f"I9{put_code}",
                institution_name="Partner University",
                ror_id="03ccc3333",
                country_code="US",
                author_count=1,
                has_corresponding_author=False,
                created_at=updated_at,
            ),
        ])
        if suffix == "a":
            db.session.add_all([
                OpenAlexWorkInstitution(
                    doi_normalized=doi,
                    openalex_id=f"W{put_code}",
                    institution_id=f"I{put_code}{index}",
                    institution_name=name,
                    ror_id=f"0{index}extra{index:04d}",
                    country_code="CL",
                    author_count=1,
                    has_corresponding_author=False,
                    created_at=updated_at,
                )
                for index, name in enumerate(
                    [
                        "Alpha Institute",
                        "Beta Institute",
                        "Delta Institute",
                        "Gamma Institute",
                        "Omega Institute",
                    ],
                    start=4,
                )
            ])
        return canonical

    def _login(
        self,
        user_id,
        *,
        is_manager=False,
        is_admin=False,
        is_oai_user=False,
        ror_id="01aaa1111",
    ):
        with self.client.session_transaction() as client_session:
            client_session.clear()
            client_session.update(
                logged_in=True,
                user_id=user_id,
                is_manager=is_manager,
                is_admin=is_admin,
                is_oai_user=is_oai_user,
                ror_id=ror_id,
                admin_selected_ror="02bbb2222",
                locale="en",
            )

    def _set_policy(self, policy):
        with self.app.app_context():
            config = OaiPmhInstitutionConfig.query.filter_by(ror_id="01aaa1111").one()
            config.publication_policy = policy
            config.policy_updated_at = utc_now().replace(microsecond=0)
            db.session.commit()

    @staticmethod
    def _doi_workbook(*values, header="doi"):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "DOI"
        worksheet.append([header])
        for value in values:
            worksheet.append([value])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        return output

    def test_standard_user_index_is_bound_to_database_institution(self):
        self._login(self.standard_id, ror_id="02bbb2222")
        configuration_context = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: configuration_context.update(context) or template_name,
        ):
            response = self.client.get("/oai-pmh/")

        self.assertEqual(200, response.status_code)
        self.assertEqual("01aaa1111", configuration_context["institution"]["ror_id"])
        self.assertEqual(2, configuration_context["stats"]["total"])
        self.assertEqual(1, configuration_context["stats"]["validated"])
        self.assertEqual(1, configuration_context["stats"]["included"])
        self.assertFalse(configuration_context["can_manage_content"])

        article_context = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: article_context.update(context) or template_name,
        ):
            article_response = self.client.get("/oai-pmh/articles/")
        self.assertEqual(200, article_response.status_code)
        self.assertTrue(all(row.ror_id == "01aaa1111" for row in article_context["records"]))
        self.assertFalse(article_context["can_manage_content"])

    def test_oai_user_can_manage_content_but_not_provider_settings(self):
        self._login(self.oai_editor_id, is_oai_user=True)
        captured = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: captured.update(context) or template_name,
        ):
            response = self.client.get("/oai-pmh/articles/?sort=title&direction=asc")

        self.assertEqual(200, response.status_code)
        self.assertTrue(captured["can_manage_content"])
        self.assertEqual(["Article A", "Article A2"], [
            record.title for record in captured["records"]
        ])

        configuration_context = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: configuration_context.update(context) or template_name,
        ):
            configuration_response = self.client.get("/oai-pmh/")
        self.assertEqual(200, configuration_response.status_code)
        self.assertFalse(configuration_context["can_manage_provider"])

        selected_record = next(
            record for record in captured["records"] if record.title == "Article A"
        )
        self.assertEqual(7, len(selected_record.affiliations))
        self.assertEqual(5, len(selected_record.displayed_affiliations))
        self.assertEqual(2, selected_record.additional_affiliation_count)
        self.assertEqual(
            ["01aaa1111"],
            [
                item["ror_id"]
                for item in selected_record.displayed_affiliations
                if item["is_selected"]
            ],
        )

        mapping_response = self.client.post(
            "/oai-pmh/metadata-mapping",
            data={
                "mapping_fields": ["title", "identifier", "volume"],
                "mapping_title": "dc.title",
                "mapping_identifier": "dc.identifier",
                "mapping_volume": "repo.volume",
            },
        )
        selection_response = self.client.post(
            "/oai-pmh/works/publication",
            data={"single_exclude": str(self.work_a_id)},
        )
        settings_response = self.client.post(
            "/oai-pmh/settings",
            data={
                "repository_name": "Unauthorized name",
                "admin_email": "other@example.org",
                "publication_policy": "all",
                "provider_enabled": "on",
            },
        )
        self.assertEqual(302, mapping_response.status_code)
        self.assertEqual(302, selection_response.status_code)
        self.assertEqual(302, settings_response.status_code)

        with self.app.app_context():
            config = OaiPmhInstitutionConfig.query.filter_by(ror_id="01aaa1111").one()
            self.assertEqual("repo.volume", config.metadata_mapping["volume"])
            self.assertEqual("DataORCID — University A", config.repository_name)
            selection = OaiPmhWorkSelection.query.filter_by(
                ror_id="01aaa1111",
                canonical_work_id=self.work_a_id,
            ).one()
            self.assertFalse(selection.is_included)

    def test_oai_user_can_download_the_guided_doi_template(self):
        self._login(self.oai_editor_id, is_oai_user=True)
        response = self.client.get("/oai-pmh/works/doi-import/template")

        self.assertEqual(200, response.status_code)
        self.assertEqual(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response.mimetype,
        )
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        try:
            self.assertIn("DOI", workbook.sheetnames)
            self.assertEqual("doi", workbook["DOI"]["A1"].value)
            self.assertIn("Instructions", workbook.sheetnames)
        finally:
            workbook.close()

    def test_article_audit_exports_include_all_filtered_institutional_rows(self):
        self._login(self.standard_id)
        csv_response = self.client.get(
            "/oai-pmh/works/export?format=csv&sort=title&direction=asc"
        )

        self.assertEqual(200, csv_response.status_code)
        self.assertEqual("text/csv", csv_response.mimetype)
        csv_rows = list(csv.reader(StringIO(csv_response.data.decode("utf-8-sig"))))
        self.assertEqual(3, len(csv_rows))
        self.assertEqual("Article", csv_rows[0][0])
        self.assertEqual(["Article A", "Article A2"], [row[0] for row in csv_rows[1:]])
        self.assertNotIn("Article B", {row[0] for row in csv_rows})
        self.assertIn("University A", csv_rows[1][6])
        self.assertIn("Omega Institute", csv_rows[1][6])

        with self.app.app_context():
            metadata = OpenAlexWorkMetadata.query.filter_by(
                doi_normalized="10.1234/a",
            ).one()
            metadata.title = "=Article A\x0b"
            db.session.commit()

        xlsx_response = self.client.get(
            "/oai-pmh/works/export?format=xlsx&validation=validated"
        )
        self.assertEqual(200, xlsx_response.status_code)
        workbook = load_workbook(BytesIO(xlsx_response.data), read_only=True)
        try:
            self.assertIn("Articles", workbook.sheetnames)
            self.assertIn("Audit summary", workbook.sheetnames)
            article_rows = list(workbook["Articles"].iter_rows(values_only=True))
            self.assertEqual(2, len(article_rows))
            self.assertEqual("'=Article A", article_rows[1][0])
            self.assertEqual("Validated", article_rows[1][9])
        finally:
            workbook.close()

    def test_doi_excel_import_is_institution_scoped_filterable_and_reversible(self):
        self._login(self.oai_editor_id, is_oai_user=True)
        with self.app.app_context():
            now = utc_now().replace(microsecond=0)
            db.session.add(OaiPmhWorkSelection(
                ror_id="01aaa1111",
                canonical_work_id=self.work_a_id,
                is_included=False,
                decision_source="manual",
                updated_by_user_id=self.oai_editor_id,
                created_at=now,
                updated_at=now,
            ))
            db.session.commit()

        response = self.client.post(
            "/oai-pmh/works/doi-import",
            data={
                "doi_file": (
                    self._doi_workbook(
                        "https://doi.org/10.1234/a",
                        "10.1234/a2",
                        "10.1234/b",
                        "not-a-doi",
                        "doi:10.1234/a",
                    ),
                    "institutional-dois.xlsx",
                ),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(302, response.status_code)
        self.assertTrue(response.location.endswith("/oai-pmh/doi-import/"))
        with self.app.app_context():
            batch = OaiPmhDoiImportBatch.query.one()
            batch_id = batch.id
            self.assertEqual(5, batch.submitted_count)
            self.assertEqual(2, batch.matched_count)
            self.assertEqual(2, batch.article_count)
            self.assertEqual(1, batch.invalid_count)
            self.assertEqual(1, batch.duplicate_count)
            self.assertEqual(1, batch.unmatched_count)
            self.assertEqual(2, OaiPmhDoiImportChange.query.count())
            selections = OaiPmhWorkSelection.query.order_by(
                OaiPmhWorkSelection.canonical_work_id
            ).all()
            self.assertEqual(
                [self.work_a_id, self.work_a2_id],
                [selection.canonical_work_id for selection in selections],
            )
            self.assertTrue(all(selection.is_included for selection in selections))
            self.assertTrue(all(selection.decision_source == "xlsx" for selection in selections))
            self.assertTrue(all(selection.doi_import_batch_id == batch.id for selection in selections))

        captured = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: captured.update(context) or template_name,
        ):
            filtered_response = self.client.get("/oai-pmh/articles/?source=xlsx")
        self.assertEqual(200, filtered_response.status_code)
        self.assertEqual("xlsx", captured["selected_source"])
        self.assertEqual(
            {"Article A", "Article A2"},
            {record.title for record in captured["records"]},
        )
        self.assertTrue(all(record.publication_source == "xlsx" for record in captured["records"]))

        history_context = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: history_context.update(context) or template_name,
        ):
            history_response = self.client.get("/oai-pmh/doi-import/")
        self.assertEqual(200, history_response.status_code)
        self.assertEqual(1, history_context["history"].total)
        self.assertEqual(batch_id, history_context["latest_active_id"])

        detail_context = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: detail_context.update(context) or template_name,
        ):
            detail_response = self.client.get(f"/oai-pmh/doi-import/{batch_id}/")
        self.assertEqual(200, detail_response.status_code)
        self.assertEqual(2, detail_context["changes"].total)
        self.assertTrue(detail_context["can_undo"])

        undo_response = self.client.post(
            f"/oai-pmh/works/doi-import/{batch_id}/undo"
        )
        self.assertEqual(302, undo_response.status_code)
        with self.app.app_context():
            restored = OaiPmhWorkSelection.query.filter_by(
                canonical_work_id=self.work_a_id,
            ).one()
            self.assertFalse(restored.is_included)
            self.assertEqual("manual", restored.decision_source)
            self.assertIsNone(restored.doi_import_batch_id)
            self.assertIsNone(OaiPmhWorkSelection.query.filter_by(
                canonical_work_id=self.work_a2_id,
            ).first())
            self.assertIsNotNone(OaiPmhDoiImportBatch.query.one().undone_at)

    def test_undo_doi_import_preserves_a_later_manual_change(self):
        self._login(self.oai_editor_id, is_oai_user=True)
        import_response = self.client.post(
            "/oai-pmh/works/doi-import",
            data={
                "doi_file": (
                    self._doi_workbook("10.1234/a"),
                    "one-doi.xlsx",
                ),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(302, import_response.status_code)

        manual_response = self.client.post(
            "/oai-pmh/works/publication",
            data={"single_exclude": str(self.work_a_id)},
        )
        self.assertEqual(302, manual_response.status_code)
        undo_response = self.client.post("/oai-pmh/works/doi-import/undo")
        self.assertEqual(302, undo_response.status_code)

        with self.app.app_context():
            selection = OaiPmhWorkSelection.query.filter_by(
                canonical_work_id=self.work_a_id,
            ).one()
            self.assertFalse(selection.is_included)
            self.assertEqual("manual", selection.decision_source)
            self.assertIsNone(selection.doi_import_batch_id)
            self.assertIsNotNone(OaiPmhDoiImportBatch.query.one().undone_at)

    def test_doi_import_without_matches_is_kept_in_the_audit_history(self):
        self._login(self.oai_editor_id, is_oai_user=True)
        response = self.client.post(
            "/oai-pmh/works/doi-import",
            data={
                "doi_file": (
                    self._doi_workbook("10.1234/b", "not-a-doi"),
                    "no-matches.xlsx",
                ),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(302, response.status_code)
        with self.app.app_context():
            batch = OaiPmhDoiImportBatch.query.one()
            self.assertEqual(2, batch.submitted_count)
            self.assertEqual(0, batch.matched_count)
            self.assertEqual(0, batch.article_count)
            self.assertEqual(1, batch.invalid_count)
            self.assertEqual(1, batch.unmatched_count)
            self.assertEqual(0, OaiPmhDoiImportChange.query.count())

        detail_context = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: detail_context.update(context) or template_name,
        ):
            detail_response = self.client.get(f"/oai-pmh/doi-import/{batch.id}/")
        self.assertEqual(200, detail_response.status_code)
        self.assertEqual(0, detail_context["changes"].total)
        self.assertTrue(detail_context["can_undo"])

    def test_standard_user_cannot_manage_oai_content(self):
        self._login(self.standard_id)
        mapping_response = self.client.post(
            "/oai-pmh/metadata-mapping",
            data={
                "mapping_title": "dc.titulo",
                "mapping_identifier": "dc.identifier",
            },
        )
        selection_response = self.client.post(
            "/oai-pmh/works/publication",
            data={"single_exclude": str(self.work_a_id)},
        )
        import_response = self.client.post(
            "/oai-pmh/works/doi-import",
            data={
                "doi_file": (
                    self._doi_workbook("10.1234/a"),
                    "unauthorized.xlsx",
                ),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(302, mapping_response.status_code)
        self.assertEqual(302, selection_response.status_code)
        self.assertEqual(302, import_response.status_code)
        with self.app.app_context():
            config = OaiPmhInstitutionConfig.query.filter_by(ror_id="01aaa1111").one()
            self.assertIsNone(config.metadata_mapping)
            self.assertEqual(0, OaiPmhWorkSelection.query.count())
            self.assertEqual(0, OaiPmhDoiImportBatch.query.count())

    def test_manager_cannot_exclude_another_institutions_work(self):
        self._login(self.manager_id, is_manager=True)
        response = self.client.post(
            "/oai-pmh/works/publication",
            data={
                "action": "exclude",
                "work_ids": [str(self.work_a_id), str(self.work_b_id)],
            },
        )

        self.assertEqual(302, response.status_code)
        with self.app.app_context():
            selections = OaiPmhWorkSelection.query.all()
            self.assertEqual(1, len(selections))
            self.assertEqual(self.work_a_id, selections[0].canonical_work_id)
            self.assertFalse(selections[0].is_included)

    def test_quick_action_changes_only_the_clicked_article(self):
        self._login(self.manager_id, is_manager=True)
        response = self.client.post(
            "/oai-pmh/works/publication",
            data={
                "action": "include",
                "work_ids": [str(self.work_a_id), str(self.work_a2_id)],
                "single_exclude": str(self.work_a_id),
            },
        )

        self.assertEqual(302, response.status_code)
        with self.app.app_context():
            selections = OaiPmhWorkSelection.query.all()
            self.assertEqual(1, len(selections))
            self.assertEqual(self.work_a_id, selections[0].canonical_work_id)
            self.assertFalse(selections[0].is_included)

    def test_default_policy_exposes_only_openalex_validated_institutional_works(self):
        records = self.client.get(
            f"/oai/{'a' * 48}?verb=ListRecords&metadataPrefix=oai_dc"
        )
        root = ET.fromstring(records.data)
        titles = [
            element.text
            for element in root.findall(".//{http://purl.org/dc/elements/1.1/}title")
        ]
        self.assertIn("Article A", titles)
        self.assertNotIn(b"Article B", records.data)
        self.assertIn(b"Author One", records.data)
        self.assertIn(b"10.1234/a", records.data)
        self.assertIn(b"Research integrity", records.data)
        self.assertNotIn(b"0.9000", records.data)
        self.assertNotIn(b"01aaa1111", records.data)

        unvalidated = self.client.get(
            f"/oai/{'a' * 48}",
            query_string={
                "verb": "GetRecord",
                "metadataPrefix": "oai_dc",
                "identifier": f"oai:dataorcid-chile:{'a' * 48}:{self.key_a2}",
            },
        )
        self.assertIn(b'code="idDoesNotExist"', unvalidated.data)
        self.assertNotIn(b"<header", unvalidated.data)
        self.assertNotIn(b"Article A2", unvalidated.data)

        identifiers = self.client.get(
            f"/oai/{'a' * 48}?verb=ListIdentifiers&metadataPrefix=oai_dc"
        )
        self.assertNotIn(b'status="deleted"', identifiers.data)
        self.assertNotIn(self.key_a2.encode(), identifiers.data)

    def test_manual_inclusion_overrides_failed_openalex_validation(self):
        self._login(self.manager_id, is_manager=True)
        self.client.post(
            "/oai-pmh/works/publication",
            data={"action": "include", "work_ids": str(self.work_a2_id)},
        )
        response = self.client.get(
            f"/oai/{'a' * 48}",
            query_string={
                "verb": "GetRecord",
                "metadataPrefix": "oai_dc",
                "identifier": f"oai:dataorcid-chile:{'a' * 48}:{self.key_a2}",
            },
        )
        self.assertIn(b"Article A2", response.data)
        self.assertNotIn(b'status="deleted"', response.data)

    def test_all_policy_can_explicitly_expose_unvalidated_articles_by_default(self):
        self._set_policy("all")
        response = self.client.get(
            f"/oai/{'a' * 48}",
            query_string={
                "verb": "GetRecord",
                "metadataPrefix": "oai_dc",
                "identifier": f"oai:dataorcid-chile:{'a' * 48}:{self.key_a2}",
            },
        )
        self.assertIn(b"Article A2", response.data)
        self.assertNotIn(b'status="deleted"', response.data)

    def test_exclusion_removes_the_article_from_the_public_repository(self):
        self._login(self.manager_id, is_manager=True)
        response = self.client.post(
            "/oai-pmh/works/publication",
            data={"action": "exclude", "work_ids": str(self.work_a_id)},
        )
        self.assertEqual(302, response.status_code)

        identifier = f"oai:dataorcid-chile:{'a' * 48}:{self.key_a}"
        record_response = self.client.get(
            f"/oai/{'a' * 48}",
            query_string={
                "verb": "GetRecord",
                "metadataPrefix": "oai_dc",
                "identifier": identifier,
            },
        )
        root = ET.fromstring(record_response.data)
        error = root.find(f".//{{{OAI_NS}}}error")
        self.assertEqual("idDoesNotExist", error.get("code"))
        self.assertIsNone(root.find(f".//{{{OAI_NS}}}header"))

        identify = self.client.get(f"/oai/{'a' * 48}?verb=Identify")
        identify_root = ET.fromstring(identify.data)
        self.assertEqual("no", identify_root.findtext(f".//{{{OAI_NS}}}deletedRecord"))

    def test_selected_only_policy_exposes_explicit_inclusions(self):
        self._set_policy("selected")
        self._login(self.manager_id, is_manager=True)
        self.client.post(
            "/oai-pmh/works/publication",
            data={"action": "include", "work_ids": str(self.work_a_id)},
        )

        included = self.client.get(
            f"/oai/{'a' * 48}",
            query_string={
                "verb": "GetRecord",
                "metadataPrefix": "oai_dc",
                "identifier": f"oai:dataorcid-chile:{'a' * 48}:{self.key_a}",
            },
        )
        excluded = self.client.get(
            f"/oai/{'a' * 48}",
            query_string={
                "verb": "GetRecord",
                "metadataPrefix": "oai_dc",
                "identifier": f"oai:dataorcid-chile:{'a' * 48}:{self.key_a2}",
            },
        )
        self.assertIn(b"Article A", included.data)
        self.assertNotIn(b'status="deleted"', included.data)
        self.assertIn(b'code="idDoesNotExist"', excluded.data)
        self.assertNotIn(b"Article A2", excluded.data)

    def test_list_records_uses_scoped_resumption_tokens(self):
        self._set_policy("all")
        with self.app.app_context():
            first_work = db.session.get(CanonicalWork, self.work_a_id)
            second_work = db.session.get(CanonicalWork, self.work_a2_id)
            shared_second = utc_now().replace(microsecond=100000)
            first_work.updated_at = shared_second
            second_work.updated_at = shared_second.replace(microsecond=900000)
            db.session.commit()

        first = self.client.get(
            f"/oai/{'a' * 48}?verb=ListIdentifiers&metadataPrefix=oai_dc"
        )
        first_root = ET.fromstring(first.data)
        token = first_root.findtext(f".//{{{OAI_NS}}}resumptionToken")
        self.assertTrue(token)
        first_identifier = first_root.findtext(
            f".//{{{OAI_NS}}}header/{{{OAI_NS}}}identifier"
        )

        second = self.client.get(
            f"/oai/{'a' * 48}",
            query_string={"verb": "ListIdentifiers", "resumptionToken": token},
        )
        second_root = ET.fromstring(second.data)
        identifiers = second_root.findall(f".//{{{OAI_NS}}}header/{{{OAI_NS}}}identifier")
        final_token = second_root.find(f".//{{{OAI_NS}}}resumptionToken")
        self.assertEqual(1, len(identifiers))
        self.assertNotEqual(first_identifier, identifiers[0].text)
        self.assertIsNotNone(final_token)
        self.assertFalse(final_token.text)

        wrong_scope = self.client.get(
            f"/oai/{'b' * 48}",
            query_string={"verb": "ListIdentifiers", "resumptionToken": token},
        )
        self.assertIn(b'code="badResumptionToken"', wrong_scope.data)

    def test_protocol_errors_remain_oai_responses(self):
        response = self.client.get(f"/oai/{'a' * 48}?verb=Unknown")
        self.assertEqual(200, response.status_code)
        self.assertIn(b'code="badVerb"', response.data)
        root = ET.fromstring(response.data)
        self.assertEqual({}, root.find(f".//{{{OAI_NS}}}request").attrib)

        empty_token = self.client.get(
            f"/oai/{'a' * 48}",
            query_string={"verb": "ListSets", "resumptionToken": ""},
        )
        self.assertIn(b'code="badResumptionToken"', empty_token.data)

    def test_provider_xml_has_a_valid_browser_presentation(self):
        response = self.client.get(f"/oai/{'a' * 48}?verb=Identify")
        self.assertIn(b'xml-stylesheet type="text/xsl"', response.data)

        xml_document = etree.fromstring(response.data)
        stylesheet = etree.parse("app/static/xsl/oai-pmh.xsl")
        html_document = etree.XSLT(stylesheet)(xml_document)
        rendered = str(html_document)
        self.assertIn("Data ORCID-Chile", rendered)
        self.assertIn("DataORCID — University A", rendered)
        self.assertIn("Identidad del repositorio", rendered)

    def test_provider_supports_form_encoded_post_requests_without_csrf(self):
        response = self.client.post(f"/oai/{'a' * 48}", data={"verb": "Identify"})
        self.assertEqual(200, response.status_code)
        self.assertIn(b"<protocolVersion>2.0</protocolVersion>", response.data)

    def test_settings_contain_no_external_source_and_can_enable_provider(self):
        self._login(self.manager_id, is_manager=True)
        response = self.client.post(
            "/oai-pmh/settings",
            data={
                "repository_name": "DataORCID University A",
                "admin_email": "admin@example.org",
                "publication_policy": "validated",
                "provider_enabled": "on",
                "source_base_url": "https://should-be-ignored.example/oai",
            },
        )
        self.assertEqual(302, response.status_code)
        with self.app.app_context():
            config = OaiPmhInstitutionConfig.query.filter_by(ror_id="01aaa1111").one()
            self.assertTrue(config.provider_enabled)
            self.assertEqual("validated", config.publication_policy)
            self.assertFalse(hasattr(config, "source_base_url"))

    def test_rotating_public_key_invalidates_the_previous_url(self):
        self._login(self.manager_id, is_manager=True)
        old_key = "a" * 48
        response = self.client.post("/oai-pmh/access-key/rotate")
        self.assertEqual(302, response.status_code)

        with self.app.app_context():
            config = OaiPmhInstitutionConfig.query.filter_by(ror_id="01aaa1111").one()
            new_key = config.public_key
        self.assertNotEqual(old_key, new_key)
        self.assertEqual(404, self.client.get(f"/oai/{old_key}?verb=Identify").status_code)
        self.assertEqual(200, self.client.get(f"/oai/{new_key}?verb=Identify").status_code)

    def test_provider_advertises_and_emits_openaire_metadata(self):
        formats = self.client.get(f"/oai/{'a' * 48}?verb=ListMetadataFormats")
        self.assertIn(b"<metadataPrefix>oai_dc</metadataPrefix>", formats.data)
        self.assertIn(b"<metadataPrefix>oai_openaire</metadataPrefix>", formats.data)
        self.assertIn(b"<metadataPrefix>dataorcid</metadataPrefix>", formats.data)
        sets = self.client.get(f"/oai/{'a' * 48}?verb=ListSets")
        self.assertIn(b"<setSpec>openaire</setSpec>", sets.data)

        response = self.client.get(
            f"/oai/{'a' * 48}?verb=ListRecords&metadataPrefix=oai_openaire&set=openaire"
        )
        root = ET.fromstring(response.data)
        datacite = "http://datacite.org/schema/kernel-4"
        oaire = "http://namespace.openaire.eu/schema/oaire/"
        self.assertEqual("Article A", root.findtext(f".//{{{datacite}}}title"))
        self.assertEqual("Author One", root.findtext(f".//{{{datacite}}}creatorName"))
        self.assertEqual("DOI", root.find(f".//{{{datacite}}}identifier").get("identifierType"))
        resource_type = root.find(f".//{{{oaire}}}resourceType")
        self.assertEqual("http://purl.org/coar/resource_type/c_6501", resource_type.get("uri"))
        self.assertNotIn(b"Article A2", response.data)

    def test_institution_can_save_and_publish_a_custom_metadata_mapping(self):
        self._login(self.manager_id, is_manager=True)
        response = self.client.post(
            "/oai-pmh/metadata-mapping",
            data={
                "mapping_title": "dc.titulo",
                "mapping_identifier": "repo.identificador",
            },
        )
        self.assertEqual(302, response.status_code)

        mapped = self.client.get(
            f"/oai/{'a' * 48}?verb=ListRecords&metadataPrefix=dataorcid"
        )
        root = ET.fromstring(mapped.data)
        namespace = "https://dataorcid.cl/ns/mapped-metadata/1.0/"
        fields = root.findall(f".//{{{namespace}}}field")
        values = {(field.get("name"), field.text) for field in fields}
        self.assertIn(("dc.titulo", "Article A"), values)
        self.assertTrue(any(name == "repo.identificador" for name, _value in values))
        self.assertFalse(any(field.get("source") == "creator" for field in fields))

        with self.app.app_context():
            config = OaiPmhInstitutionConfig.query.filter_by(ror_id="01aaa1111").one()
            self.assertEqual("dc.titulo", config.metadata_mapping["title"])

    def test_metadata_mapping_is_a_separate_view_with_optional_source_fields(self):
        self._login(self.manager_id, is_manager=True)
        captured = {}
        with patch(
            "app.blueprints.oai_pmh.render_template",
            side_effect=lambda template_name, **context: captured.update(context) or template_name,
        ):
            response = self.client.get("/oai-pmh/metadata/")

        self.assertEqual(200, response.status_code)
        self.assertEqual("oai_pmh/metadata_mapping.html", response.text)
        catalog = {field["key"]: field for field in captured["metadata_catalog"]}
        self.assertTrue(catalog["title"]["active"])
        self.assertFalse(catalog["volume"]["active"])
        self.assertFalse(catalog["institution_ror"]["active"])
        self.assertIn("metadataPrefix=oai_dc", captured["format_urls"]["oai_dc"])
        self.assertIn(
            "metadataPrefix=oai_openaire",
            captured["format_urls"]["oai_openaire"],
        )
        self.assertIn(
            "metadataPrefix=dataorcid",
            captured["format_urls"]["dataorcid"],
        )

    def test_custom_mapping_can_add_extended_fields_and_hide_default_fields(self):
        self._login(self.manager_id, is_manager=True)
        response = self.client.post(
            "/oai-pmh/metadata-mapping",
            data={
                "mapping_fields": ["title", "identifier", "volume", "institution_ror"],
                "mapping_title": "dc.title",
                "mapping_identifier": "dc.identifier",
                "mapping_volume": "repo.volume",
                "mapping_institution_ror": "repo.institution.ror",
            },
        )
        self.assertEqual(302, response.status_code)
        self.assertTrue(response.location.endswith("/oai-pmh/metadata/"))

        mapped = self.client.get(
            f"/oai/{'a' * 48}?verb=ListRecords&metadataPrefix=dataorcid"
        )
        root = ET.fromstring(mapped.data)
        namespace = "https://dataorcid.cl/ns/mapped-metadata/1.0/"
        fields = root.findall(f".//{{{namespace}}}field")
        by_source = {}
        for field in fields:
            by_source.setdefault(field.get("source"), []).append(
                (field.get("name"), field.text)
            )
        self.assertIn(("repo.volume", "42"), by_source["volume"])
        self.assertIn(
            ("repo.institution.ror", "https://ror.org/01aaa1111"),
            by_source["institution_ror"],
        )
        self.assertNotIn("creator", by_source)

    def test_custom_mapping_rejects_invalid_destination_names(self):
        self._login(self.manager_id, is_manager=True)
        self.client.post(
            "/oai-pmh/metadata-mapping",
            data={
                "mapping_title": "<titulo>",
                "mapping_identifier": "dc.identifier",
            },
        )
        with self.app.app_context():
            config = OaiPmhInstitutionConfig.query.filter_by(ror_id="01aaa1111").one()
            self.assertIsNone(config.metadata_mapping)

    def test_repository_summary_uses_effective_exposure_counts(self):
        with self.app.app_context():
            summaries = oai_repository_summaries()
            by_ror = {item["config"].ror_id: item for item in summaries}

            self.assertEqual(2, by_ror["01aaa1111"]["available"])
            self.assertEqual(1, by_ror["01aaa1111"]["validated"])
            self.assertEqual(1, by_ror["01aaa1111"]["exposed"])
            self.assertEqual(1, by_ror["02bbb2222"]["available"])

    def test_oai_administration_is_scoped_for_managers_and_global_for_admins(self):
        manager_context = {}
        self._login(self.manager_id, is_manager=True)
        with patch(
            "app.blueprints.admin.render_template",
            side_effect=lambda template_name, **context: manager_context.update(context) or template_name,
        ):
            response = self.client.get("/admin/oai-pmh")
        self.assertEqual(200, response.status_code)
        self.assertEqual(["01aaa1111"], [
            item["config"].ror_id for item in manager_context["repositories"]
        ])

        admin_context = {}
        self._login(self.admin_id, is_admin=True)
        with patch(
            "app.blueprints.admin.render_template",
            side_effect=lambda template_name, **context: admin_context.update(context) or template_name,
        ):
            response = self.client.get("/admin/oai-pmh")
        self.assertEqual(200, response.status_code)
        self.assertEqual(2, admin_context["summary"]["repositories"])

    def test_admin_can_open_an_institutional_oai_workspace(self):
        self._login(self.admin_id, is_admin=True)
        response = self.client.post("/admin/oai-pmh/01aaa1111/open")
        self.assertEqual(302, response.status_code)
        self.assertTrue(response.location.endswith("/oai-pmh/"))
        with self.client.session_transaction() as client_session:
            self.assertEqual("01aaa1111", client_session["admin_selected_ror"])


if __name__ == "__main__":
    unittest.main()
