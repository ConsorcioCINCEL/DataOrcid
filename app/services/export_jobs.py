"""Durable, access-controlled generation of large CSV and XLSX exports."""

from __future__ import annotations

import csv
import hashlib
import json
import logging
import os
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from threading import Lock
from typing import Iterable

import pandas as pd
from flask import abort, current_app, jsonify, send_file, session, url_for
from sqlalchemy import and_, func, or_

from .. import db
from ..models import SyncJob, SyncJobStep, utc_now
from ..spreadsheet import excel_safe_dataframe
from .background_jobs import submit_background_job, update_job_progress

logger = logging.getLogger(__name__)
_CLEANUP_LOCK = Lock()
_LAST_CLEANUP_MONOTONIC = 0.0
_EXPORT_FORMATS = {"csv", "excel"}
_ACTIVE_EXPORT_STATUSES = {"queued", "running"}
_TERMINAL_EXPORT_STATUSES = {"success", "failed", "partial", "interrupted"}
# Bump this value whenever the generated columns or formatting change so an
# older, otherwise valid file is never reused after an export-format release.
_EXPORT_CONTENT_VERSION = 1
_EXPORT_MIMETYPES = {
    "csv": "text/csv; charset=utf-8",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _export_directory() -> Path:
    directory = Path(
        current_app.config.get("EXPORT_DIRECTORY")
        or (Path(current_app.instance_path) / "exports")
    )
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    return directory


def _job_artifact_paths(job: SyncJob, root: Path) -> set[Path]:
    """Return only paths that are safely contained in the private export root."""
    paths: set[Path] = set()
    result = job.result_json if isinstance(job.result_json, dict) else {}
    export = result.get("export") if isinstance(result.get("export"), dict) else {}
    stored_name = export.get("stored_name")
    if isinstance(stored_name, str) and stored_name and Path(stored_name).name == stored_name:
        candidate = root / stored_name
        if candidate.parent == root:
            paths.add(candidate)

    job_id = str(job.id or "")
    if job_id and Path(job_id).name == job_id:
        for extension in ("csv", "xlsx"):
            paths.add(root / f"{job_id}.{extension}")
            paths.add(root / f".{job_id}.{extension}.tmp")
    return paths


def _remove_job_artifacts(job: SyncJob, root: Path, removed: set[Path]) -> None:
    for path in _job_artifact_paths(job, root):
        if path in removed or not path.is_file():
            continue
        try:
            path.unlink()
            removed.add(path)
        except FileNotFoundError:
            continue
        except OSError:
            logger.warning("Could not remove generated export file %s.", path, exc_info=True)


def _delete_export_job_rows(job_ids: list[str]) -> int:
    if not job_ids:
        return 0
    SyncJobStep.query.filter(SyncJobStep.sync_job_id.in_(job_ids)).delete(
        synchronize_session=False
    )
    deleted = SyncJob.query.filter(
        SyncJob.id.in_(job_ids),
        SyncJob.job_type == "export",
        SyncJob.status.in_(_TERMINAL_EXPORT_STATUSES),
    ).delete(synchronize_session=False)
    db.session.commit()
    return deleted


def cleanup_expired_exports(*, force: bool = False) -> dict[str, int]:
    """Remove expired export artifacts and their terminal job history."""
    global _LAST_CLEANUP_MONOTONIC
    now_monotonic = time.monotonic()
    with _CLEANUP_LOCK:
        if not force and now_monotonic - _LAST_CLEANUP_MONOTONIC < 300:
            return {"files": 0, "jobs": 0}
        _LAST_CLEANUP_MONOTONIC = now_monotonic

        cutoff = utc_now() - timedelta(
            hours=max(int(current_app.config.get("EXPORT_RETENTION_HOURS", 24)), 1)
        )
        root = _export_directory().resolve()
        removed_files: set[Path] = set()
        removed_jobs = 0
        expired_condition = or_(
            and_(SyncJob.finished_at.isnot(None), SyncJob.finished_at < cutoff),
            and_(SyncJob.finished_at.is_(None), SyncJob.created_at < cutoff),
        )

        try:
            expired_jobs = SyncJob.query.filter(
                SyncJob.job_type == "export",
                SyncJob.status.in_(_TERMINAL_EXPORT_STATUSES),
                expired_condition,
            ).all()
            for job in expired_jobs:
                _remove_job_artifacts(job, root, removed_files)
            removed_jobs = _delete_export_job_rows([job.id for job in expired_jobs])
        except Exception:
            db.session.rollback()
            logger.warning("Could not prune expired export job history.", exc_info=True)

        try:
            protected_ids = {
                job_id
                for (job_id,) in db.session.query(SyncJob.id).filter(
                    SyncJob.job_type == "export",
                )
            }
            for path in root.iterdir():
                try:
                    if not path.is_file() or path in removed_files:
                        continue
                    if any(
                        path.name.startswith(f"{job_id}.")
                        or path.name.startswith(f".{job_id}.")
                        for job_id in protected_ids
                    ):
                        continue
                    modified_at = datetime.fromtimestamp(
                        path.stat().st_mtime,
                        tz=timezone.utc,
                    ).replace(tzinfo=None)
                    if modified_at < cutoff:
                        path.unlink()
                        removed_files.add(path)
                except FileNotFoundError:
                    continue
                except OSError:
                    logger.warning(
                        "Could not remove orphaned export file %s.",
                        path,
                        exc_info=True,
                    )
        except OSError:
            logger.warning("Could not inspect the private export directory.", exc_info=True)

        return {"files": len(removed_files), "jobs": removed_jobs}


def _safe_format(value: str | None) -> str:
    return "excel" if (value or "").lower() in {"excel", "xlsx"} else "csv"


def _datetime_revision(value) -> str:
    return value.isoformat(timespec="microseconds") if value else ""


def _table_revision(model, timestamp_column, *criteria) -> list:
    query = db.session.query(func.count(), func.max(timestamp_column))
    if criteria:
        query = query.filter(*criteria)
    count, latest = query.one()
    return [int(count or 0), _datetime_revision(latest)]


def _export_source_state(kind: str, parameters: dict) -> tuple[str, datetime | None]:
    """Build a revision and latest change time for an export's source data."""
    from ..models import (
        DuplicateProfileReview,
        FundingCacheRun,
        InstitutionRegistry,
        InstitutionResearcher,
        OaiPmhDoiImportBatch,
        OaiPmhInstitutionConfig,
        OaiPmhWorkSelection,
        OpenAlexSyncRun,
        OpenAlexWorkMetadata,
        ResearcherCache,
        ResearcherStatus,
        WorkCacheRun,
    )

    ror_id = parameters.get("ror_id")
    revisions: dict[str, object] = {"content_version": _EXPORT_CONTENT_VERSION}

    source_job_query = db.session.query(
        func.count(SyncJob.id),
        func.max(SyncJob.created_at),
        func.max(SyncJob.finished_at),
    ).filter(SyncJob.job_type != "export")
    source_count, source_created, source_finished = source_job_query.one()
    revisions["source_jobs"] = [
        int(source_count or 0),
        _datetime_revision(source_created),
        _datetime_revision(source_finished),
    ]

    work_kinds = {
        "institution_works",
        "staff_works",
        "institution_researchers",
        "staff_researchers",
        "researcher_directory",
        "institution_summary",
        "raw_openalex",
        "institution_openalex",
        "duplicate_profiles",
        "oai_article_audit",
    }
    funding_kinds = {
        "institution_fundings",
        "staff_fundings",
        "institution_researchers",
        "staff_researchers",
        "researcher_directory",
        "institution_summary",
        "duplicate_profiles",
    }
    researcher_kinds = {
        "institution_researchers",
        "staff_researchers",
        "researcher_directory",
        "institution_summary",
        "duplicate_profiles",
    }
    openalex_kinds = {
        "raw_openalex",
        "institution_openalex",
        "institution_summary",
        "oai_article_audit",
    }

    if kind in work_kinds:
        criteria = (WorkCacheRun.ror_id == ror_id,) if ror_id else ()
        revisions["work_runs"] = _table_revision(
            WorkCacheRun,
            WorkCacheRun.finished_at,
            *criteria,
        )
    if kind in funding_kinds:
        criteria = (FundingCacheRun.ror_id == ror_id,) if ror_id else ()
        revisions["funding_runs"] = _table_revision(
            FundingCacheRun,
            FundingCacheRun.finished_at,
            *criteria,
        )
    if kind in researcher_kinds:
        status_criteria = (ResearcherStatus.ror_id == ror_id,) if ror_id else ()
        revisions["researcher_status"] = _table_revision(
            ResearcherStatus,
            ResearcherStatus.last_updated,
            *status_criteria,
        )
        revisions["researcher_profiles"] = _table_revision(
            ResearcherCache,
            ResearcherCache.updated_at,
        )
        association_query = db.session.query(
            func.count(InstitutionResearcher.id),
            func.max(InstitutionResearcher.last_seen_at),
            func.max(InstitutionResearcher.profile_updated_at),
        ).join(
            InstitutionRegistry,
            InstitutionRegistry.id == InstitutionResearcher.institution_id,
        )
        if ror_id:
            association_query = association_query.filter(
                InstitutionRegistry.ror_id == ror_id
            )
        association_count, last_seen, profile_updated = association_query.one()
        revisions["associations"] = [
            int(association_count or 0),
            _datetime_revision(last_seen),
            _datetime_revision(profile_updated),
        ]
        registry_criteria = (InstitutionRegistry.ror_id == ror_id,) if ror_id else ()
        revisions["institutions"] = _table_revision(
            InstitutionRegistry,
            InstitutionRegistry.updated_at,
            *registry_criteria,
        )
    if kind in openalex_kinds:
        run_criteria = (
            or_(OpenAlexSyncRun.ror_id == ror_id, OpenAlexSyncRun.ror_id.is_(None)),
        ) if ror_id else ()
        revisions["openalex_runs"] = _table_revision(
            OpenAlexSyncRun,
            OpenAlexSyncRun.finished_at,
            *run_criteria,
        )
        revisions["openalex_metadata"] = _table_revision(
            OpenAlexWorkMetadata,
            OpenAlexWorkMetadata.updated_at,
        )
    if kind == "duplicate_profiles":
        ror_ids = [value for value in (parameters.get("ror_ids") or []) if value]
        review_criteria = (
            (DuplicateProfileReview.ror_id.in_(ror_ids),) if ror_ids else ()
        )
        revisions["duplicate_reviews"] = _table_revision(
            DuplicateProfileReview,
            DuplicateProfileReview.updated_at,
            *review_criteria,
        )
    if kind == "oai_article_audit":
        revisions["oai_config"] = _table_revision(
            OaiPmhInstitutionConfig,
            OaiPmhInstitutionConfig.updated_at,
            OaiPmhInstitutionConfig.ror_id == ror_id,
        )
        revisions["oai_selections"] = _table_revision(
            OaiPmhWorkSelection,
            OaiPmhWorkSelection.updated_at,
            OaiPmhWorkSelection.ror_id == ror_id,
        )
        batch_count, batch_created, batch_undone = db.session.query(
            func.count(OaiPmhDoiImportBatch.id),
            func.max(OaiPmhDoiImportBatch.created_at),
            func.max(OaiPmhDoiImportBatch.undone_at),
        ).filter(OaiPmhDoiImportBatch.ror_id == ror_id).one()
        revisions["oai_imports"] = [
            int(batch_count or 0),
            _datetime_revision(batch_created),
            _datetime_revision(batch_undone),
        ]

    encoded = json.dumps(
        revisions,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    timestamps: list[datetime] = []

    def collect_timestamps(value) -> None:
        if isinstance(value, dict):
            for nested in value.values():
                collect_timestamps(nested)
        elif isinstance(value, (list, tuple)):
            for nested in value:
                collect_timestamps(nested)
        elif isinstance(value, str) and value:
            try:
                timestamps.append(datetime.fromisoformat(value))
            except ValueError:
                pass

    collect_timestamps(revisions)
    return (
        hashlib.sha256(encoded).hexdigest()[:20],
        max(timestamps) if timestamps else None,
    )


def _export_source_revision(kind: str, parameters: dict) -> str:
    return _export_source_state(kind, parameters)[0]


def _reusable_export_path(job: SyncJob) -> Path | None:
    if job.status != "success":
        return None
    result = job.result_json if isinstance(job.result_json, dict) else {}
    export = result.get("export") if isinstance(result.get("export"), dict) else {}
    stored_name = export.get("stored_name")
    if not isinstance(stored_name, str) or Path(stored_name).name != stored_name:
        return None
    retention = timedelta(
        hours=max(int(current_app.config.get("EXPORT_RETENTION_HOURS", 24)), 1)
    )
    completed_at = job.finished_at or job.created_at
    if not completed_at or completed_at + retention < utc_now():
        return None
    root = _export_directory().resolve()
    path = (root / stored_name).resolve()
    return path if path.parent == root and path.is_file() else None


def _dataframe_export(
    data_frame: pd.DataFrame,
    temporary_path: Path,
    export_format: str,
    sheet_name: str,
) -> int:
    if export_format == "excel":
        excel_frame = excel_safe_dataframe(data_frame)
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            excel_frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    else:
        data_frame.to_csv(temporary_path, index=False, encoding="utf-8-sig")
    return len(data_frame.index)


def _cache_dataframe(kind: str, parameters: dict) -> tuple[pd.DataFrame, str, str]:
    from ..blueprints.works_shared import (
        _build_fundings_dataframe,
        _build_researchers_dataframe,
        _build_works_dataframe,
        _format_datetime,
        _institution_cache_summaries,
    )
    from ..models import FundingCache, WorkCache

    ror_id = parameters.get("ror_id")
    if kind == "institution_works":
        rows = WorkCache.query.filter_by(ror_id=ror_id).all()
        return _build_works_dataframe(rows), f"orcid_works_cache_{ror_id}", "Works"
    if kind == "institution_fundings":
        rows = FundingCache.query.filter_by(ror_id=ror_id).all()
        return _build_fundings_dataframe(rows), f"orcid_fundings_cache_{ror_id}", "Fundings"
    if kind == "institution_researchers":
        return (
            _build_researchers_dataframe(ror_id),
            f"orcid_researchers_{ror_id}",
            "Researchers",
        )
    if kind == "staff_works":
        rows = WorkCache.query.order_by(WorkCache.ror_id, WorkCache.orcid, WorkCache.id).all()
        return _build_works_dataframe(rows), "orcid_works_all_institutions", "Works"
    if kind == "staff_fundings":
        rows = FundingCache.query.order_by(FundingCache.ror_id, FundingCache.orcid, FundingCache.id).all()
        return _build_fundings_dataframe(rows), "orcid_fundings_all_institutions", "Fundings"
    if kind == "staff_researchers":
        return _build_researchers_dataframe(), "orcid_researchers_all_institutions", "Researchers"
    if kind == "institution_summary":
        rows = [{
            "institution": item["name"],
            "ror_id": item["ror_id"],
            "researchers": item["researchers"],
            "works": item["works"],
            "fundings": item["fundings"],
            "openalex_matched": item["openalex_matched"],
            "openalex_candidates": item["openalex_candidates"],
            "openalex_coverage_percent": item["openalex_percent"],
            "last_update": _format_datetime(item["last_update"]),
            "cache_status": item["health"],
        } for item in _institution_cache_summaries()]
        return pd.DataFrame(rows), "institution_cache_summary", "Institution summary"
    if kind == "researcher_directory":
        from flask_babel import force_locale
        from ..blueprints.main import _researcher_directory_export_dataframe

        with force_locale(parameters.get("locale") or "en"):
            data_frame = _researcher_directory_export_dataframe(
                ror_id,
                parameters.get("filters") or {},
                parameters.get("institution_name"),
            )
        return data_frame, f"researchers_{ror_id}", "Researchers"
    raise ValueError(f"Unsupported dataframe export kind: {kind}")


def _write_dict_rows(
    rows: Iterable[dict],
    columns: list[str],
    temporary_path: Path,
    export_format: str,
    sheet_name: str,
    job_id: str | None,
) -> int:
    count = 0
    if export_format == "excel":
        from openpyxl import Workbook
        from ..blueprints.works_shared import _excel_cell

        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet(sheet_name[:31])
        worksheet.append(columns)
        worksheet_rows = 1
        sheet_number = 1
        for row in rows:
            if worksheet_rows >= 1_048_576:
                sheet_number += 1
                worksheet = workbook.create_sheet(f"{sheet_name[:24]} {sheet_number}")
                worksheet.append(columns)
                worksheet_rows = 1
            worksheet.append([_excel_cell(row.get(column)) for column in columns])
            worksheet_rows += 1
            count += 1
            if count % 500 == 0:
                update_job_progress(job_id, count, 0, "records")
        workbook.save(temporary_path)
    else:
        with temporary_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow(row)
                count += 1
                if count % 500 == 0:
                    update_job_progress(job_id, count, 0, "records")
    return count


def _raw_openalex_rows(query):
    from ..blueprints.works_shared import _openalex_export_row

    for raw, metadata in query.yield_per(500):
        yield _openalex_export_row(raw, metadata)


def _raw_openalex_query(parameters: dict):
    from ..blueprints.works_shared import (
        _openalex_cache_key_expr,
        _openalex_export_load_options,
    )
    from ..models import OpenAlexWorkMetadata, OpenAlexWorkRawCache, WorkCache

    ror_id = parameters.get("ror_id")
    if ror_id:
        cache_key = _openalex_cache_key_expr(WorkCache).label("openalex_cache_key")
        cache_keys = (
            db.session.query(cache_key)
            .filter(WorkCache.ror_id == ror_id, WorkCache.type == "journal-article")
            .distinct()
            .subquery()
        )
        query = (
            db.session.query(OpenAlexWorkRawCache, OpenAlexWorkMetadata)
            .join(cache_keys, OpenAlexWorkRawCache.doi_normalized == cache_keys.c.openalex_cache_key)
            .outerjoin(
                OpenAlexWorkMetadata,
                OpenAlexWorkMetadata.doi_normalized == OpenAlexWorkRawCache.doi_normalized,
            )
            .order_by(OpenAlexWorkRawCache.doi_normalized.asc())
        )
    else:
        query = (
            db.session.query(OpenAlexWorkRawCache, OpenAlexWorkMetadata)
            .outerjoin(
                OpenAlexWorkMetadata,
                OpenAlexWorkMetadata.doi_normalized == OpenAlexWorkRawCache.doi_normalized,
            )
            .order_by(OpenAlexWorkRawCache.doi_normalized.asc())
        )
    return query.options(
        *_openalex_export_load_options(OpenAlexWorkRawCache, OpenAlexWorkMetadata)
    )


def _write_raw_openalex(
    parameters: dict,
    temporary_path: Path,
    export_format: str,
    job_id: str | None,
) -> tuple[int, str, str]:
    from ..blueprints.works_shared import (
        OPENALEX_EXPORT_CSV_COLUMNS,
        OPENALEX_EXPORT_XLSX_COLUMNS,
    )

    ror_id = parameters.get("ror_id")
    columns = OPENALEX_EXPORT_XLSX_COLUMNS if export_format == "excel" else OPENALEX_EXPORT_CSV_COLUMNS
    count = _write_dict_rows(
        _raw_openalex_rows(_raw_openalex_query(parameters)),
        columns,
        temporary_path,
        export_format,
        "OpenAlex",
        job_id,
    )
    base_name = f"openalex_articles_{ror_id}" if ror_id else "openalex_articles_all_institutions"
    return count, base_name, "OpenAlex"


def _write_institution_openalex(
    parameters: dict,
    temporary_path: Path,
    export_format: str,
    job_id: str | None,
) -> tuple[int, str, str]:
    from ..blueprints.works_openalex_data import (
        OPENALEX_INSTITUTION_EXPORT_COLUMNS,
        _openalex_institution_export_query,
        _openalex_institution_export_row,
    )

    ror_id = parameters["ror_id"]
    coverage = parameters.get("coverage") or "all"
    query = _openalex_institution_export_query(
        ror_id,
        coverage=coverage,
        search=parameters.get("search") or "",
        sort=parameters.get("sort") or "citations",
        direction=parameters.get("direction") or "desc",
    )
    rows = (
        _openalex_institution_export_row(work, raw, metadata)
        for work, raw, metadata in query.yield_per(1000)
    )
    count = _write_dict_rows(
        rows,
        OPENALEX_INSTITUTION_EXPORT_COLUMNS,
        temporary_path,
        export_format,
        "OpenAlex works",
        job_id,
    )
    return count, f"openalex_works_{ror_id}_{coverage}", "OpenAlex works"


def _write_duplicate_profiles(
    parameters: dict,
    temporary_path: Path,
    export_format: str,
) -> tuple[int, str, str]:
    from flask_babel import _, force_locale
    from ..blueprints.duplicates import (
        _activity_column_labels,
        _activity_column_order,
        _export_column_labels,
        _export_column_order,
        _institution_column_labels,
        _localized_export_rows,
        _methodology_rows,
    )
    from .duplicate_profile_service import build_duplicate_report, flatten_duplicate_rows

    with force_locale(parameters.get("locale") or "en"):
        report = build_duplicate_report(ror_ids=parameters.get("ror_ids"))
        rows = _localized_export_rows(flatten_duplicate_rows(report["groups"]))
        candidates = pd.DataFrame(rows, columns=_export_column_order())
        candidates = candidates.rename(columns=_export_column_labels())
        if export_format == "excel":
            with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
                excel_safe_dataframe(candidates).to_excel(
                    writer, sheet_name=_("Candidates"), index=False
                )
                institutions = pd.DataFrame(report["institutions"]).rename(
                    columns=_institution_column_labels()
                )
                excel_safe_dataframe(institutions).to_excel(
                    writer, sheet_name=_("Institutions"), index=False
                )
                activity = pd.DataFrame(
                    report["profile_activity"], columns=_activity_column_order()
                ).rename(columns=_activity_column_labels())
                excel_safe_dataframe(activity).to_excel(
                    writer,
                    sheet_name=_("ORCID Activity"),
                    index=False,
                )
                excel_safe_dataframe(pd.DataFrame(_methodology_rows())).to_excel(
                    writer,
                    sheet_name=_("Methodology"),
                    index=False,
                )
        else:
            candidates.to_csv(temporary_path, index=False, encoding="utf-8-sig")
    scope = "all" if parameters.get("scope") == "all" else "current"
    return len(rows), f"duplicate-orcid-profiles-{scope}", "Candidates"


def _write_oai_article_audit(
    parameters: dict,
    temporary_path: Path,
    export_format: str,
    job_id: str | None,
) -> tuple[int, str, str]:
    import re

    from flask_babel import _, force_locale
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill

    from ..blueprints.oai_pmh import (
        _article_export_headers,
        _article_listing,
        _iter_article_export_rows,
        _safe_spreadsheet_cell,
    )
    from ..models import CanonicalWork, OaiPmhInstitutionConfig, utc_now

    ror_id = parameters["ror_id"]
    institution_name = parameters.get("institution_name") or ror_id
    with force_locale(parameters.get("locale") or "en"):
        stored_config = OaiPmhInstitutionConfig.query.filter_by(ror_id=ror_id).first()
        config = stored_config or OaiPmhInstitutionConfig(
            ror_id=ror_id,
            repository_name=f"DataORCID-Chile — {institution_name}",
            public_key="",
            provider_enabled=False,
            publication_policy="validated",
            policy_updated_at=utc_now().replace(microsecond=0),
            created_at=utc_now().replace(microsecond=0),
            updated_at=utc_now().replace(microsecond=0),
        )
        listing = _article_listing(
            config,
            parameters.get("filters") or {},
            load_type_options=False,
        )
        ordered_query = listing["query"].order_by(listing["ordering"], CanonicalWork.id.desc())
        headers = _article_export_headers()
        rows = _iter_article_export_rows(ordered_query, config, ror_id)
        exported_count = 0

        if export_format == "csv":
            with temporary_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                for row in rows:
                    writer.writerow([_safe_spreadsheet_cell(value) for value in row])
                    exported_count += 1
                    if exported_count % 500 == 0:
                        update_job_progress(job_id, exported_count, 0, "records")
        else:
            workbook = Workbook(write_only=True)
            worksheet = workbook.create_sheet(_("Articles"))
            worksheet.freeze_panes = "A2"
            for column, width in zip(
                "ABCDEFGHIJKLMN",
                [48, 28, 45, 10, 18, 30, 60, 45, 38, 20, 16, 22, 22, 48],
            ):
                worksheet.column_dimensions[column].width = width
            header_fill = PatternFill("solid", fgColor="C94725")
            header_cells = []
            for value in headers:
                cell = WriteOnlyCell(worksheet, value=value)
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center")
                header_cells.append(cell)
            worksheet.append(header_cells)
            for row in rows:
                worksheet.append([
                    _safe_spreadsheet_cell(value, xlsx=True) for value in row
                ])
                exported_count += 1
                if exported_count % 500 == 0:
                    update_job_progress(job_id, exported_count, 0, "records")
            worksheet.auto_filter.ref = f"A1:N{max(exported_count + 1, 1)}"

            summary = workbook.create_sheet(_("Audit summary"))
            summary.column_dimensions["A"].width = 30
            summary.column_dimensions["B"].width = 70
            summary_rows = [
                [_("Institution"), institution_name],
                ["ROR", ror_id],
                [_("Generated at"), utc_now().strftime("%Y-%m-%dT%H:%M:%SZ")],
                [_("Exported articles"), exported_count],
                [_("Search filter"), listing["query_text"] or _("None")],
                [_("OAI status filter"), listing["status"]],
                [_("OpenAlex validation filter"), listing["validation"]],
                [_("Activation-origin filter"), listing["source"]],
            ]
            for summary_row in summary_rows:
                summary.append([
                    _safe_spreadsheet_cell(value, xlsx=True) for value in summary_row
                ])
            workbook.save(temporary_path)
            workbook.close()

    timestamp = utc_now().strftime("%Y%m%d_%H%M%S")
    safe_ror = re.sub(r"[^A-Za-z0-9_-]", "_", ror_id)
    return exported_count, f"oai_articulos_{safe_ror}_{timestamp}", "Articles"


def generate_export_file(
    kind: str,
    export_format: str,
    parameters: dict,
    label: str,
    job_id: str | None = None,
) -> dict:
    """Generate one export into private storage and return public-safe metadata."""
    export_format = _safe_format(export_format)
    extension = "xlsx" if export_format == "excel" else "csv"
    directory = _export_directory()
    stored_name = f"{job_id}.{extension}"
    final_path = directory / stored_name
    temporary_path = directory / f".{job_id}.{extension}.tmp"
    temporary_path.unlink(missing_ok=True)

    try:
        if kind in {
            "institution_works",
            "institution_fundings",
            "institution_researchers",
            "institution_summary",
            "staff_works",
            "staff_fundings",
            "staff_researchers",
            "researcher_directory",
        }:
            data_frame, base_name, sheet_name = _cache_dataframe(kind, parameters)
            count = _dataframe_export(data_frame, temporary_path, export_format, sheet_name)
        elif kind == "raw_openalex":
            count, base_name, sheet_name = _write_raw_openalex(
                parameters, temporary_path, export_format, job_id
            )
        elif kind == "institution_openalex":
            count, base_name, sheet_name = _write_institution_openalex(
                parameters, temporary_path, export_format, job_id
            )
        elif kind == "duplicate_profiles":
            count, base_name, sheet_name = _write_duplicate_profiles(
                parameters,
                temporary_path,
                export_format,
            )
        elif kind == "oai_article_audit":
            count, base_name, sheet_name = _write_oai_article_audit(
                parameters,
                temporary_path,
                export_format,
                job_id,
            )
        else:
            raise ValueError(f"Unsupported export kind: {kind}")

        os.replace(temporary_path, final_path)
        try:
            final_path.chmod(0o600)
        except OSError:
            logger.debug("Could not adjust export file permissions for %s", final_path)
        update_job_progress(job_id, count, count, "records", message=f"{label}: {count} records.")
        return {
            "export": {
                "stored_name": stored_name,
                "filename": f"{base_name}.{extension}",
                "mimetype": _EXPORT_MIMETYPES[export_format],
                "size_bytes": final_path.stat().st_size,
                "records": count,
                "format": export_format,
                "label": label,
            }
        }
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def _job_label(job: SyncJob) -> str:
    try:
        args = (job.payload_json or {}).get("args") or []
        return str(args[3]) if len(args) > 3 else job.name
    except (TypeError, ValueError):
        return job.name


def _can_access_export(job: SyncJob) -> bool:
    return bool(
        job.job_type == "export"
        and (
            job.requested_by_user_id == session.get("user_id")
            or session.get("is_admin")
        )
    )


def export_job_payload(job: SyncJob) -> dict:
    """Serialize an export job without exposing its private filesystem path."""
    result = job.result_json if isinstance(job.result_json, dict) else {}
    export = result.get("export") if isinstance(result.get("export"), dict) else {}
    ready = job.status == "success" and bool(export.get("stored_name"))
    payload = {
        "id": job.id,
        "label": export.get("label") or _job_label(job),
        "status": job.status,
        "message": job.message,
        "created_at": job.created_at.isoformat() if job.created_at else None,
        "finished_at": job.finished_at.isoformat() if job.finished_at else None,
        "current": int(job.items_current or 0),
        "total": int(job.items_total or 0),
        "filename": export.get("filename"),
        "size_bytes": export.get("size_bytes"),
        "records": export.get("records"),
        "download_url": (
            url_for("background_exports.download_export", job_id=job.id)
            if ready
            else None
        ),
    }
    if job.status in {"failed", "partial", "interrupted"}:
        payload["error"] = "The export could not be completed."
    return payload


def _export_job_module_enabled(job: SyncJob) -> bool:
    """Return whether the module that owns an export is currently available."""
    from .module_access import is_module_enabled, module_for_export_job

    return is_module_enabled(module_for_export_job(job))


def _disabled_export_module_response(job: SyncJob):
    from .module_access import module_disabled_response, module_for_export_job

    module_key = module_for_export_job(job)
    return module_disabled_response(module_key) if module_key else None


def queue_export_response(
    kind: str,
    export_format: str,
    parameters: dict,
    label: str,
):
    """Queue an export for the current account and return an HTTP 202 payload."""
    from .module_access import (
        is_module_enabled,
        module_disabled_response,
        module_for_export_kind,
    )

    module_key = module_for_export_kind(kind)
    if module_key and not is_module_enabled(module_key):
        return module_disabled_response(module_key)
    cleanup_expired_exports()
    export_format = _safe_format(export_format)
    normalized_parameters = json.loads(json.dumps(parameters))
    legacy_signature = hashlib.sha256(
        json.dumps(normalized_parameters, sort_keys=True).encode("utf-8")
    ).hexdigest()[:16]
    source_revision, source_changed_at = _export_source_state(
        kind,
        normalized_parameters,
    )
    normalized_parameters["_source_revision"] = source_revision
    signature = hashlib.sha256(
        json.dumps(normalized_parameters, sort_keys=True).encode("utf-8")
    ).hexdigest()[:16]
    user_id = int(session["user_id"])
    job_name = f"export-{user_id}-{kind}-{export_format}-{signature}"
    legacy_job_name = (
        f"export-{user_id}-{kind}-{export_format}-{legacy_signature}"
    )
    candidates = (
        SyncJob.query
        .filter_by(
            name=job_name,
            job_type="export",
            requested_by_user_id=user_id,
        )
        .filter(SyncJob.status.in_(_ACTIVE_EXPORT_STATUSES | {"success"}))
        .order_by(SyncJob.created_at.desc())
        .all()
    )
    if legacy_job_name != job_name:
        candidates.extend(
            SyncJob.query
            .filter_by(
                name=legacy_job_name,
                job_type="export",
                requested_by_user_id=user_id,
                status="success",
            )
            .order_by(SyncJob.created_at.desc())
            .all()
        )
    existing = next(
        (
            job
            for job in candidates
            if job.name == job_name and job.status in _ACTIVE_EXPORT_STATUSES
        ),
        None,
    )
    reusable_paths = {
        job.id: _reusable_export_path(job)
        for job in candidates
        if job.status == "success"
    }
    current_successes = [
        job
        for job in candidates
        if job.name == job_name and reusable_paths.get(job.id) is not None
    ]
    legacy_successes = [
        job
        for job in candidates
        if (
            job.name == legacy_job_name
            and reusable_paths.get(job.id) is not None
            and (
                source_changed_at is None
                or (job.finished_at or job.created_at) >= source_changed_at
            )
        )
    ]
    reusable = next(iter(current_successes), None) or next(
        iter(legacy_successes),
        None,
    )
    stale_successes = [
        job
        for job in candidates
        if job.status == "success"
        and (
            reusable_paths.get(job.id) is None
            or (job.name == job_name and job is not reusable)
        )
    ]
    if stale_successes:
        root = _export_directory().resolve()
        removed_files: set[Path] = set()
        for job in stale_successes:
            _remove_job_artifacts(job, root, removed_files)
        _delete_export_job_rows([job.id for job in stale_successes])
    if reusable:
        if reusable.name == legacy_job_name:
            reusable.name = job_name
            payload = json.loads(json.dumps(reusable.payload_json or {}))
            args = payload.get("args") if isinstance(payload, dict) else None
            if isinstance(args, list) and len(args) > 2 and isinstance(args[2], dict):
                args[2] = normalized_parameters
                reusable.payload_json = payload
            db.session.commit()
        return jsonify({"job": export_job_payload(reusable), "reused": True}), 202

    active_query = SyncJob.query.filter_by(job_type="export").filter(
        SyncJob.status.in_(_ACTIVE_EXPORT_STATUSES)
    )
    if not existing:
        user_active = active_query.filter_by(requested_by_user_id=user_id).count()
        global_active = active_query.count()
        if user_active >= int(current_app.config.get("EXPORT_MAX_ACTIVE_PER_USER", 3)):
            return jsonify({"error": "Too many active exports for this account."}), 429
        if global_active >= int(current_app.config.get("EXPORT_MAX_ACTIVE_GLOBAL", 20)):
            return jsonify({"error": "The export queue is currently full."}), 503
    job_id = submit_background_job(
        current_app._get_current_object(),
        job_name,
        generate_export_file,
        kind,
        export_format,
        normalized_parameters,
        label,
        job_type="export",
        ror_id=normalized_parameters.get("ror_id"),
        requested_by_user_id=user_id,
        steps=["prepare_file"],
        deduplicate=True,
    )
    job = db.session.get(SyncJob, job_id)
    return jsonify({
        "job": export_job_payload(job),
        "reused": existing is not None,
    }), 202


def list_current_user_exports():
    """Return recent export jobs so notifications survive page navigation."""
    cleanup_expired_exports()
    cutoff = utc_now() - timedelta(
        hours=max(int(current_app.config.get("EXPORT_RETENTION_HOURS", 24)), 1)
    )
    jobs = (
        SyncJob.query
        .filter_by(job_type="export", requested_by_user_id=session.get("user_id"))
        .filter(or_(
            SyncJob.status.in_(_ACTIVE_EXPORT_STATUSES),
            and_(SyncJob.finished_at.isnot(None), SyncJob.finished_at >= cutoff),
            and_(SyncJob.finished_at.is_(None), SyncJob.created_at >= cutoff),
        ))
        .order_by(SyncJob.created_at.desc())
        .limit(12)
        .all()
    )
    return jsonify({
        "jobs": [
            export_job_payload(job)
            for job in jobs
            if _export_job_module_enabled(job)
        ]
    })


def clear_current_user_exports():
    """Delete every finished export owned by the current account."""
    jobs = SyncJob.query.filter(
        SyncJob.job_type == "export",
        SyncJob.requested_by_user_id == int(session["user_id"]),
        SyncJob.status.in_(_TERMINAL_EXPORT_STATUSES),
    ).all()
    root = _export_directory().resolve()
    removed_files: set[Path] = set()
    for job in jobs:
        _remove_job_artifacts(job, root, removed_files)
    removed_jobs = _delete_export_job_rows([job.id for job in jobs])
    return jsonify({
        "deleted_files": len(removed_files),
        "deleted_jobs": removed_jobs,
    })


def get_export_job_response(job_id: str):
    job = db.session.get(SyncJob, job_id)
    if not job or not _can_access_export(job):
        abort(404)
    if not _export_job_module_enabled(job):
        return _disabled_export_module_response(job)
    return jsonify({"job": export_job_payload(job)})


def download_export_response(job_id: str):
    job = db.session.get(SyncJob, job_id)
    if not job or not _can_access_export(job):
        abort(404)
    if not _export_job_module_enabled(job):
        return _disabled_export_module_response(job)
    result = job.result_json if isinstance(job.result_json, dict) else {}
    export = result.get("export") if isinstance(result.get("export"), dict) else {}
    stored_name = export.get("stored_name")
    if job.status != "success" or not stored_name or Path(stored_name).name != stored_name:
        abort(404)

    root = _export_directory().resolve()
    path = (root / stored_name).resolve()
    if path.parent != root or not path.is_file():
        abort(410)
    retention = timedelta(
        hours=max(int(current_app.config.get("EXPORT_RETENTION_HOURS", 24)), 1)
    )
    if job.finished_at and job.finished_at + retention < utc_now():
        path.unlink(missing_ok=True)
        abort(410)
    response = send_file(
        path,
        as_attachment=True,
        download_name=export.get("filename") or stored_name,
        mimetype=export.get("mimetype") or "application/octet-stream",
        conditional=True,
        max_age=0,
    )
    response.headers["Cache-Control"] = "private, no-store"
    return response
